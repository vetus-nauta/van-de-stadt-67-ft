"""Independent workbook fixtures, synthetic money ONLY under .local/U01/review."""
from pathlib import Path
import json, shutil
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[5]; LOCAL=ROOT/'.local/U01/review'; SRC=ROOT/'.local/U01/UTV-template-draft2.xlsx'
OUT=LOCAL/'input';OUT.mkdir(parents=True,exist_ok=True)
shutil.copy2(SRC,OUT/'blank.xlsx')
w=load_workbook(SRC)
def setv(s,r,data):
 for c,v in data.items():w[s][f'{c}{r}']=v
# All original template slots remain empty except synthetic scenario rows below.
for row in w['07_Смета'].iter_rows(min_row=4):
 for c in row:
  if c.data_type!='f':c.value=None
setv('01_Основания',8,{'B':'EUR'});setv('01_Основания',9,{'B':'NET'});setv('01_Основания',10,{'B':'2026-09-10'})
setv('13_Источники',4,{'A':'SRC-TEST','B':'SYNTHETIC','C':'Synthetic fixture source','D':'2026-09-10','E':'R1','F':'.local/U01/review/TEST-ONLY','H':'PRIVATE_LOCAL','I':'DOCUMENT_VERIFIED'})
cases=[]
def fixture(name,r,kind='KAC',expected='READY',amount=200):
 q=f'Q{r}';line=f'QL{r}';package=f'RFQ{r}';option=f'O{r}'
 setv('03_Пакеты',r,{'A':package,'B':'W00.10','C':name,'D':'R1','E':'Synthetic technical brief','G':'pc','O':'SRC-TEST','Q':'READY_TO_REQUEST'})
 setv('04_КП',r,{'A':q,'B':'SUPPLIER_TEST','C':'TEST','D':'2026-09-10','E':'2026-10-10','F':'EUR','G':'NET','H':'SYNTHETIC TAX BASIS','I':0,'M':'INCLUDED','N':'INCLUDED','P':'SRC-TEST'})
 setv('05_Строки_КП',r,{'A':line,'B':q,'C':package,'D':'line1','E':'SYNTHETIC ITEM','F':2,'G':'pc','H':100,'I':'KNOWN','N':'SRC-TEST/page1'})
 setv('06_КАЦ',r,{'A':option,'B':f'GROUP{r}','C':line,'D':package,'E':'COMPARABLE','F':'pc','G':1,'H':1,'I':'2026-09-10','J':'SRC-TEST','K':0,'L':0,'M':0,'N':0,'S':'SELECTED','T':'SYNTHETIC TEST','U':'Same unit; zero extras confirmed synthetic','W':'SRC-TEST'})
 setv('06а_Оценки',r,{'A':f'D{r}','B':package,'C':'CALCULATION','D':'SYNTHETIC TEST','E':'pc','F':100,'G':'EUR','H':'NET','I':'2026-09-10','J':'SRC-TEST','K':'Synthetic calculation','L':'KNOWN','M':'TEST DECISION'})
 setv('07_Смета',r,{'A':f'EST-TEST-{r}','B':'W00.10','C':package,'D':name,'E':'EQUIPMENT','F':'NEW','G':'OWNER','H':'REMAINING','I':2,'J':'pc','K':option if kind=='KAC' else f'D{r}','O':'SRC-TEST','P':'R1/page1','Q':'KNOWN','R':'Explicit synthetic scope','V':kind})
 cases.append({'name':name,'row':r,'expected_status':expected,'expected_amount':amount})
fixture('valid_kac',4)
fixture('valid_direct',5,'DIRECT')
fixture('confirmed_zero_price',6,amount=0);setv('05_Строки_КП',6,{'H':0,'I':'ZERO_CONFIRMED'})
fixture('unknown_zero_price',7,expected='INCOMPLETE / LINK',amount=None);setv('05_Строки_КП',7,{'H':0,'I':'UNKNOWN'})
fixture('missing_price',8,expected='INCOMPLETE / LINK',amount=None);setv('05_Строки_КП',8,{'H':None})
fixture('missing_fx',9,expected='INCOMPLETE / LINK',amount=None);setv('06_КАЦ',9,{'H':None})
fixture('missing_tax_basis',10,expected='INCOMPLETE / LINK',amount=None);setv('04_КП',10,{'G':None})
fixture('technical_reject',11,expected='INCOMPLETE / LINK',amount=None);setv('06_КАЦ',11,{'E':'NOT_COMPARABLE'})
fixture('double_selected_a',12,expected='INCOMPLETE / LINK',amount=None)
fixture('double_selected_b',13,expected='INCOMPLETE / LINK',amount=None);setv('06_КАЦ',13,{'B':'GROUP12'})
fixture('unit_mismatch',14,expected='INCOMPLETE / LINK',amount=None);setv('07_Смета',14,{'J':'kg'})
fixture('package_mismatch',15,expected='INCOMPLETE / LINK',amount=None);setv('06_КАЦ',15,{'D':'RFQ4'})
fixture('negative_quantity',16,expected='INCOMPLETE / LINK',amount=None);setv('07_Смета',16,{'I':-1})
fixture('negative_price',17,expected='INCOMPLETE / LINK',amount=None);setv('05_Строки_КП',17,{'H':-1})
fixture('rounding_kac',18,amount=3.02);setv('05_Строки_КП',18,{'H':1.005});setv('07_Смета',18,{'I':3})
fixture('rounding_direct',19,'DIRECT',amount=3.02);setv('06а_Оценки',19,{'F':1.005});setv('07_Смета',19,{'I':3})
fixture('sunk_existing',20,amount=200);setv('07_Смета',20,{'H':'SUNK','F':'EXISTING','G':'OWNER'})
fixture('confirmed_zero_quantity',21,amount=0);setv('07_Смета',21,{'I':0,'Q':'ZERO_CONFIRMED'})
fixture('missing_quantity',22,expected='INCOMPLETE / LINK',amount=None);setv('07_Смета',22,{'I':None})
fixture('missing_quantity_revision',23,expected='INCOMPLETE / LINK',amount=None);setv('07_Смета',23,{'P':None})
fixture('missing_correction',24,expected='INCOMPLETE / LINK',amount=None);setv('06_КАЦ',24,{'M':None})
fixture('unresolved_tax_other',25,expected='INCOMPLETE / LINK',amount=None);setv('04_КП',25,{'G':'OTHER','H':None,'I':None})
fixture('missing_quote_price_date',26,expected='INCOMPLETE / LINK',amount=None);setv('04_КП',26,{'D':None})
fixture('missing_source_registry',27,expected='INCOMPLETE / LINK',amount=None);setv('04_КП',27,{'P':'DOES-NOT-EXIST'})
fixture('unconfirmed_zero_total_by_adjustment',28,expected='INCOMPLETE / LINK',amount=None);setv('06_КАЦ',28,{'N':-100})
fixture('not_applicable',29,expected='NOT_APPLICABLE',amount=None);setv('07_Смета',29,{'Q':'NOT_APPLICABLE','H':'OUTSIDE','I':None,'K':None})
fixture('invalid_numeric_text',30,expected='INCOMPLETE / LINK',amount=None);setv('07_Смета',30,{'I':'two'})
fixture('direct_negative',31,'DIRECT',expected='INCOMPLETE / LINK',amount=None);setv('06а_Оценки',31,{'F':-1})
fixture('direct_tax_mismatch',32,'DIRECT',expected='INCOMPLETE / LINK',amount=None);setv('06а_Оценки',32,{'H':'GROSS'})
fixture('unit_fx_adjustment_arithmetic',33,amount=2420);setv('06_КАЦ',33,{'G':3,'H':4,'K':1,'L':2,'M':3,'N':4})
fixture('unselected_option',34,expected='INCOMPLETE / LINK',amount=None);setv('06_КАЦ',34,{'S':'PENDING'})
fixture('missing_quantity_source',35,expected='INCOMPLETE / LINK',amount=None);setv('07_Смета',35,{'O':None})
fixture('missing_direct_source_id',36,'DIRECT',expected='INCOMPLETE / LINK',amount=None);setv('06а_Оценки',36,{'J':'DOES-NOT-EXIST'})
fixture('missing_fx_source_id',37,expected='INCOMPLETE / LINK',amount=None);setv('06_КАЦ',37,{'J':'DOES-NOT-EXIST'})
fixture('missing_adjustment_source_id',38,expected='INCOMPLETE / LINK',amount=None);setv('06_КАЦ',38,{'W':'DOES-NOT-EXIST'})
fixture('missing_quantity_source_id',39,expected='INCOMPLETE / LINK',amount=None);setv('07_Смета',39,{'O':'DOES-NOT-EXIST'})
setv('01_Основания',16,{'B':'TEST APPROVAL'})
for r,payment in [(4,30),(5,3000)]:setv('11_Контроль',r,{'A':'2026-09','B':'W00','C':'TEST-BASE','D':1000,'E':500,'F':200,'G':payment,'H':300,'I':400,'J':50,'M':'SYNTHETIC NONOVERLAP BASIS'})
w.save(OUT/'cases.xlsx')
setv('01_Основания',10,{'B':None});w.save(OUT/'missing-base-date.xlsx')
setv('01_Основания',10,{'B':'2026-09-10'});setv('01_Основания',9,{'B':None});w.save(OUT/'missing-base-tax.xlsx')
(LOCAL/'expected.json').write_text(json.dumps(cases,indent=2)+'\n')
print('Created four copies; original unchanged. Synthetic entries are not project prices.')
