"""Build the M06 scope schedule; no prices or approved design are inferred.

Run with .local/utv-tools/venv/bin/python from repository root.
The XLSX contains values and blank supplier response cells, no cost formulas.
"""
import csv
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

S = Path(__file__).resolve().parent
OUT = S.parent / '00_powerplant-scope.xlsx'

def read(name):
    with (S / name).open(newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f, delimiter=';'))

rows = []
for name in ['mechanical.csv', 'electrical-hydraulic.csv', 'generator-interfaces.csv']:
    rows.extend(read(name))
assert len(rows) == len({r['item_id'] for r in rows})
with (S / 'scope-register.csv').open('w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter=';', lineterminator='\n')
    w.writeheader()
    w.writerows(rows)

wb = Workbook()
wb.remove(wb.active)
navy = '183749'
blue = 'E7F1F6'
yellow = 'FFF2CC'

def sheet(title, headers, values, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in values:
        ws.append(row)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 80
    for cell in ws[1]:
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=navy)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 36
    for idx, row in enumerate(ws.iter_rows(min_row=2), 2):
        ws.row_dimensions[idx].height = 100 if title == '01_Комплектность' else 78
        for cell in row:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if idx % 2 == 0:
                cell.fill = PatternFill('solid', fgColor=blue)
            if isinstance(cell.value, str) and cell.value.startswith('https://'):
                # A multi-source note remains text; never link it as one malformed URL.
                if ' ' not in cell.value and '|' not in cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(name='Calibri', size=11, color='0563C1', underline='single')
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
    return ws

intro = [
    ['Назначение', 'Ведомость комплектности одной главной силовой линии и её интерфейсов; основа запроса поставщикам. Не утверждённый подбор, не полный бюджет, не проверенная помещаемость.'],
    ['Как пользоваться', '01: состав/границы; 02: открытые вопросы; 03: данные для размещения; 04: пустой ответ поставщика по тем же ID. Черновик письма: RFQ-draft-en.md.'],
    ['Количество', '1 главный двигатель, 1 редуктор для одной линии, 2 генератора. 6 сервисных + 2 генераторных + 1 пусковая АКБ = 9 физических; EH01 — ссылка, не повторная закупка.'],
    ['Альтернативы', '6LPA bobtail + KMH50A-LP и 8LV320 + KMH50A + VC20 — два способа собрать одну главную линию. Столбцы не складываются. Генераторы: два агрегата одного согласованного варианта, не две пары сразу.'],
    ['Электрика', 'Требование владельца:24ВDC/220ВAC. Штатные12В исполнения не признаны соответствующими; требуется решение OEM. Единица115ампер/напряжение АКБ остаются открытыми.'],
    ['Гидравлика', 'Все подруливающие устройства и лебёдки гидравлические. PTO/потребители/одновременность/резерв/тепловой баланс — отдельные позиции и интерфейсы.'],
    ['Включение', 'INCLUDED_DEALER_STATEMENT: дилер заявил включение, нужно BOM. OPTIONAL: опция, не заказ. SEPARATE: выделить внешнюю поставку/функцию. OPEN: включение/решение не установлено. N_A: не применяется или строка-ссылка.'],
    ['Количества внутри пакета', 'Пустое количество означает отсутствие расчёта; это не ноль. Сборочный комплект не подменяет число шлангов, клапанов, метров кабеля или нормочасов.'],
    ['Источники цен', 'M03 — рабочая КАЦ; M04/M05 — ценовые дополнения. Ссылки на цену детали не закрывают цену монтажа или всего узла. Palmer4544USD имеет конфликт срока цены2020; при новом КП заменить с сохранением следа изменения.'],
    ['Ответ поставщика', 'Одна копия листа04 на одного поставщика и один вариант. INCLUDED: назвать родительский пакет. Нулевую цену указывать только как явно включённую стоимость; неизвестное оставить пустым. Цены NET и VAT отдельно.'],
    ['Закрытие', 'Каждая позиция содержит требуемое доказательство, ответственного и будущий критерий приёмки. OPEN не означает выполненный расчёт или испытание.'],
    ['Область', 'Генерация, батареи/щиты, общие системы и гидравлика показаны как интерфейсы своих пакетов M03, их не закупать повторно. Полный перечень всего МО остаётся M01.'],
    ['Отправка', 'Письмо подготовлено на английском, но не отправлено. Заказы и обязательства не оформлялись.'],
    ['Версия', 'M06/v001 · 11.09.2026. Предыдущие выпуски сохраняются. Исходные размеры МО не подтверждены натурными обмерами.']
]
sheet('00_Порядок', ['Тема', 'Правило / результат'], intro, [28, 125])

headers = ['ID','Группа','Элемент / функция','Количество','Ед.','Основание количества','8LV: включение','6LPA: включение','Граница поставки','Интерфейсы','Цена: ссылка / ID','Технический источник','Что получить для закрытия','Кто закрывает','Статус','Критерий приёмки — выполнить','Примечание']
data = []
for r in rows:
    vals = list(r.values())
    vals[3] = float(r['quantity']) if r['quantity'] else None
    data.append(vals)
sheet('01_Комплектность', headers, data, [12,23,48,12,17,42,25,25,58,40,36,45,65,33,22,65,55])
questions = read('open-questions.csv')
sheet('02_Открытые вопросы', ['ID','Приоритет','Вопрос','Требуемый результат','Ответственный','Пакеты','Статус'], [list(r.values()) for r in questions], [12,17,38,90,55,25,32])
spaces = read('space-inputs.csv')
sd = []
for r in spaces:
    vals = list(r.values())
    for i in range(2, 7):
        vals[i] = float(vals[i]) if vals[i] else None
    sd.append(vals)
sheet('03_Для размещения', ['ID','Агрегат / альтернативный вариант','Количество варианта','Длина мм','Ширина мм','Высота мм','Сухая масса одного комплекта кг','Граница размеров / массы','Основание','Источник','Что требуется дальше','Статус'], sd, [12,44,17,15,15,15,22,90,40,50,70,40])

response_headers = ['ID M06','Элемент','Поставщик','№ КП / ревизия','Вариант','Включение','Родительский пакет / ID','Точный артикул','Кол-во в КП','Ед.','Цена NET / ед.','Валюта','VAT / основание','Условия / отправка','Срок поставки','Срок цены','Документ / лист','Примечание / исключения']
response = [[r['item_id'],r['item'],'','','','','','','','','','','','','','','','Строка-ссылка, не повторная закупка' if r['group']=='EXISTING_MAIN_COMPONENT_REFS' else ''] for r in rows]
ws = sheet('04_Ответ поставщика', response_headers, response, [12,48,28,25,22,26,28,28,15,15,21,14,38,45,24,24,42,60])
status = DataValidation(type='list', formula1='"INCLUDED,SEPARATE,EXCLUDED,OPTIONAL,CLARIFY,REFERENCE_ONLY"', allow_blank=True)
status.errorTitle = 'Выберите статус'
status.error = 'Используйте список статусов включения.'
status.showErrorMessage = True
ws.add_data_validation(status)
status.add(f'F2:F{ws.max_row}')
for col in ['I','K']:
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{ws.max_row}')
for row in ws.iter_rows(min_row=2, min_col=3):
    for c in row:
        c.fill = PatternFill('solid', fgColor=yellow)
for row in range(2, ws.max_row+1):
    ws.cell(row,11).number_format = '#,##0.00'

wb.save(OUT)
loaded = load_workbook(OUT, data_only=False)
assert loaded['01_Комплектность'].max_row == len(rows)+1
assert loaded['04_Ответ поставщика'].max_row == len(rows)+1
assert not any(c.data_type == 'f' for w in loaded for row in w for c in row)
stats = {'scope_rows':len(rows),'questions':len(questions),'placement_inputs':len(spaces),'worksheets':len(loaded.sheetnames),'formulas':0,'supplier_prices':'BLANK','status':'SCOPE_PREPARED_TECHNICAL_CLOSURE_OPEN'}
(S/'build-stats.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(stats,ensure_ascii=False))
