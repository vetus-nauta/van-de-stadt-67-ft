"""Build a blank, versioned yacht baseline-budget workbook. No prices or quantities."""
from pathlib import Path
import argparse,csv,json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side,Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as L
ROOT=Path(__file__).resolve().parents[5]
p=argparse.ArgumentParser();p.add_argument('--out',required=True,type=Path);a=p.parse_args()
if a.out.exists():p.error('Refusing to overwrite an existing workbook')
N=203
wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=124519,fullCalcOnLoad=True)
NAVY='19394A';TEAL='DCEDEC';BLUE='EAF2FA';GOLD='FFF2CC';GREY='EDF0F2'
schema={}
def sheet(name,title,headers,n=N):
 ws=wb.create_sheet(name);ws.append([title]);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers));ws.append(['Синие ячейки — ввод; серые — формулы. Пусто = нет данных. Код строки сохранять при сортировке.']);ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers));ws.append(headers)
 for r in range(4,n+1):
  for c in range(1,len(headers)+1):
   q=ws.cell(r,c);q.fill=PatternFill('solid',fgColor=BLUE);q.font=Font(name='Calibri',size=11,color='164B78');q.alignment=Alignment(vertical='top',wrap_text=True);q.protection=Protection(locked=False)
 for r in [1,3]:
  for cell in ws[r]:cell.fill=PatternFill('solid',fgColor=NAVY);cell.font=Font(name='Calibri',size=12 if r==3 else 18,bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=34;ws.row_dimensions[2].height=30;ws.row_dimensions[3].height=48
 for c,h in enumerate(headers,1):ws.column_dimensions[L(c)].width=22 if c<4 else 28
 ws.freeze_panes='D4';ws.auto_filter.ref=f'A3:{L(len(headers))}{n}';ws.sheet_view.showGridLines=False
 ws.print_title_rows='1:3';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_options.horizontalCentered=True;ws.oddFooter.center.text='VDS67 • U01 • TEMPLATE — не утверждённый бюджет';ws.oddFooter.right.text='Стр. &P / &N'
 tab=Table(displayName='T'+str(len(wb.worksheets)),ref=f'A3:{L(len(headers))}{n}');tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False);ws.add_table(tab)
 schema[name]={'headers':headers,'first_data_row':4,'last_data_row':n,'formula_columns':[]}
 return ws
def formula(ws,col,r,text,money=False):
 q=ws[f'{col}{r}'];q.value=text;q.fill=PatternFill('solid',fgColor=GREY);q.font=Font(name='Calibri',size=11,color='33434B');q.protection=Protection(locked=True)
 if money:q.number_format='#,##0.00;[Red](#,##0.00);0.00'
 if col not in schema[ws.title]['formula_columns']:schema[ws.title]['formula_columns'].append(col)
def choices(ws,col,values):
 d=DataValidation(type='list',formula1='"'+','.join(values)+'"',allow_blank=True);d.errorTitle='Выберите значение';d.error='Используйте значение из списка.';d.showErrorMessage=True;d.errorStyle='stop';ws.add_data_validation(d);d.add(f'{col}4:{col}{N}')
def number(ws,col,positive=False):
 d=DataValidation(type='decimal',operator='greaterThan' if positive else 'greaterThanOrEqual',formula1=0,allow_blank=True);d.showErrorMessage=True;d.error='Требуется допустимое число.';ws.add_data_validation(d);d.add(f'{col}4:{col}{N}')
def lookup(sheet,key,col):return f'IFERROR(INDEX(\'{sheet}\'!${col}$4:${col}${N},MATCH({key},\'{sheet}\'!$A$4:$A${N},0)),"")'
def count(sheet,col,key):return f'COUNTIF(\'{sheet}\'!${col}$4:${col}${N},{key})'
start=sheet('00_Старт','VDS67 • Шаблон будущей УТВ',['Раздел','Порядок работы / правило'],20)
intro=[('Статус','TEMPLATE / INCOMPLETE. УТВ возникает после отдельного решения владельца о базовом бюджете.'),('1. Основания','Установить границу остатка достройки, валюту, налоговую основу, дату цен и ревизию проекта.'),('2. Состав','WBS — системы яхты. Строки сметы разворачивать по ресурсам/работам; один пакет может содержать несколько строк.'),('3. Пакеты','Сначала составить одинаковые задания RFQ с границами поставки, монтажа и испытаний.'),('4. КП','Регистрировать документы и исходные строки без изменения оригиналов.'),('5. КАЦ','Сначала техническая сопоставимость, затем единицы, валюта, налоги и исключённые работы. Автоматического выбора минимальной цены нет.'),('6. Смета','Цена из выбранной КАЦ (KAC) или обоснованной оценки без КП (DIRECT). Количество и источник отдельно; выбор КАЦ единственный в группе.'),('7. Сводка','Показывает только известную часть и незаполненные позиции. Пустые суммы не равны нулю. Полнота состава отдельно проверяется.'),('Существующее','Наличие, поставщик и граница затрат — разные поля. Имеющийся корпус не считать новой покупкой по умолчанию.'),('Ноль','Числовой 0 допустим с подтверждением; UNKNOWN и NOT_APPLICABLE не вводить в денежные клетки.'),('Двойной счёт','Включённые монтаж, доставка и испытания не добавлять повторно. Вести лист Стыки и ссылку на покрывающую строку.'),('Резерв','Только по риску и решению. Утверждённый резерв включать отдельной строкой сметы один раз; лист Риски не прибавляется автоматически.'),('Версия','Исходный U01 неизменен. Для заполнения создать рабочую копию; при утверждении базового бюджета отдельный выпуск.'),('Конфиденциальность','Репозиторий публичный. Частные КП и контактные данные хранить в .local или согласованном закрытом хранилище; сюда безопасный ID.'),('Вместимость формы','По200 строк на регистрах. При расширении продлить формулы, проверки, диапазоны сводки и поисков; не писать за последней строкой.'),('Расчёт','Суммы в одной базовой валюте и основе; округление строк до0.01. Пустые вводные не заменяются нулями.'),('Модель','EXP-003 — принятая стартовая идея. Производственные объёмы, мощность оборудования и размеры из неё не утверждены.')]
for r,row in enumerate(intro,4):
 for c,v in enumerate(row,1):start.cell(r,c,v)
start.column_dimensions['A'].width=25;start.column_dimensions['B'].width=115
for r in range(4,21):start.row_dimensions[r].height=40
basis=sheet('01_Основания','Основания будущего бюджета',['Поле','Значение','Источник / решение'],22)
basisrows=[('Проект','VDS67','office/project.json'),('Стройка','Польша','Решение владельца'),('Статус','TEMPLATE / INCOMPLETE','Утверждение бюджета отсутствует'),('Граница бюджета','Остаток достройки; понесённое отдельно','Уточнить перечень имущества и ранее оплаченного'),('Базовая валюта',None,None),('Налоговая основа NET/GROSS',None,None),('Дата цен',None,None),('Ревизия проекта',None,None),('Владелец / заказчик',None,None),('Верфь / место передачи',None,None),('Налоговый режим / обоснование',None,None),('Курс: базовая валюта за1 валюты КП',None,None),('Решение об утверждении УТВ',None,None),('Дата утверждения',None,None),('Полнота состава проверена',None,None),('Правило округления','Строка0.01; сумма округлённых строк','Проверить договорные требования'),('Версия шаблона','U01/v001','Пустой инструмент; для заполнения копия'),('Рабочая копия / ответственный',None,None),('Флаг / эксплуатация / применимость',None,None)]
for r,row in enumerate(basisrows,4):
 for c,v in enumerate(row,1):basis.cell(r,c,v)
basis.column_dimensions['A'].width=43;basis.column_dimensions['B'].width=48;basis.column_dimensions['C'].width=70
# B8 currency; B9 tax basis; B10 date; B16 approval; B18 completeness.
d=DataValidation(type='list',formula1='"NET,GROSS"',allow_blank=True);basis.add_data_validation(d);d.add('B9')
raw=list(csv.DictReader((ROOT/'data/office/wbs.csv').open(),delimiter=';'))
packages=[]
for n in ['hull-packages.json','systems-packages.json']:
 packages += json.loads((ROOT/'deliverables/U01/v001/01_support/basis'/n).read_text())
packages += [dict(code_hint='W00.10',title='Управление и координация проекта',scope='Планирование, ведение бюджета и документов, координация исполнителей',exclusions='Техническое проектирование W02; включённые накладные подрядчиков',needed_inputs='Организация проекта; состав услуг и договоры')]
packages += [
 dict(code_hint='W02.40',title='Применимость требований и согласования',scope='Маршрут оценки соответствия, технический файл и необходимые внешние проверки по установленному флагу и эксплуатации',exclusions='Согласования не назначены автоматически; инженерная разработка в W02.10–30',needed_inputs='Статус корпуса; флаг; режим эксплуатации; требования компетентных организаций'),
 dict(code_hint='W90.40',title='Общая логистика и внешние расходы',scope='Перевозка оборудования/материалов, общая доставка и оформление по применимости',exclusions='Доставка уже включённая в цены КП или распределённая по позициям; двойное начисление запрещено',needed_inputs='Место передачи; условия КП; состав грузов; схема распределения'),
 dict(code_hint='W90.50',title='Страхование строительства и согласованные сборы',scope='Страхование рисков строительства/перевозки и проектные сборы по подтверждённым условиям',exclusions='Включённые страховки подрядчиков; эксплуатационные расходы после передачи',needed_inputs='Граница ответственности; договоры и основания сборов'),
 dict(code_hint='W80.40',title='Первичное снабжение, запасные части и инструменты',scope='Первичный комплект ЗИП, специальные инструменты, маркировка и согласованное снабжение при передаче',exclusions='Штатные комплекты уже включённые в поставку систем; последующие эксплуатационные закупки',needed_inputs='BOM выбранного оборудования; комплектность поставок; перечень сдачи')]
packages += [
 dict(code_hint='W10.85',title='Аноды и наружные элементы противокоррозионной защиты',scope='Аноды, их корпусные крепления и установка по проекту защиты алюминиевого корпуса',exclusions='Электрические защитные соединения и развязка W40.25; покрытия W10.70–80',needed_inputs='Проект защиты; материалы смежных узлов; условия эксплуатации'),
 dict(code_hint='W10.90',title='Интегральные сварные танки — условный пакет',scope='Изготовление встроенных танков, корпусные стенки/люки и испытание герметичности до подключения системы',exclusions='Покупные съёмные ёмкости и трубная обвязка W50; повторное начисление испытаний всей системы',needed_inputs='Решение о типе и материале танков; объёмы; чертежи; программа испытаний'),
 dict(code_hint='W50.58',title='Санитарные приборы и их установка',scope='Унитазы, раковины, смесители, душевые приборы и установочные комплекты; приёмка подключения по интерфейсу',exclusions='Магистрали и насосы W50.20/50; мебель и гидроизоляция помещений W60.50; бытовые приборы W60.60',needed_inputs='План санитарных зон; спецификация приборов; граница комплектов'),
 dict(code_hint='W50.85',title='Газовая система — только при выборе газового оборудования',scope='Газовый шкаф, баллоны, редуцирование, трубопроводы, контроль утечки и испытания по отдельному решению',exclusions='Полностью электрическое исполнение не принято автоматически; газовые потребители W60.60; не считать пакет обязательным',needed_inputs='Выбор газ/электричество; требования эксплуатации; проект системы')]
assert len({q['code_hint'] for q in packages})==len(packages),'Duplicate WBS'
packages.sort(key=lambda q:q['code_hint'])
wbs=sheet('02_WBS','Структура работ и систем яхты',['WBS','Родитель','Раздел / пакет','Входит','Граница / исключено','Нужные исходные данные','Применимость','Ответственный'],len(raw)+len(packages)+3)
r=4
for parent in raw:
 vals=[parent['wbs_id'],'',parent['name'],parent['scope'],'См. дочерние пакеты','Установить состав и ревизию','REVIEW_REQUIRED','']
 for c,v in enumerate(vals,1):wbs.cell(r,c,v)
 r+=1
 for pck in [q for q in packages if q['code_hint'].split('.')[0]==parent['wbs_id']]:
  vals=[pck['code_hint'],parent['wbs_id'],pck['title'],pck['scope'],pck['exclusions'],pck['needed_inputs'],'REVIEW_REQUIRED','']
  for c,v in enumerate(vals,1):wbs.cell(r,c,v)
  wbs.row_dimensions[r].height=70;r+=1
for c in ['C','D','E','F']:wbs.column_dimensions[c].width=48
choices(wbs,'G',['REVIEW_REQUIRED','IN_SCOPE','OPTIONAL','NOT_APPLICABLE'])
rfq=sheet('03_Пакеты','Задания на запрос КП / RFQ',['Пакет_ID','WBS','Название','Ревизия ТЗ','Описание и характеристики','Количество','Единица','Поставка / состав','Монтаж / изготовление','Испытания / документы','Исключено','Место передачи','Желаемый срок','Гарантия / сервис','Источник ТЗ','Ответственный','Статус'])
choices(rfq,'Q',['DRAFT','READY_TO_REQUEST','REQUESTED','OFFERS_RECEIVED','CLOSED']);number(rfq,'F')
quotes=sheet('04_КП','Реестр исходных коммерческих предложений',['КП_ID','Поставщик_ID / имя','Документ №','Дата цены','Действует до','Валюта КП','Основа NET/GROSS/OTHER','Налог / режим текстом','Ставка по документу','Условия оплаты','Срок поставки','Место / условия передачи','Доставка включена?','Монтаж включён?','Гарантия / сервис','Документ_ID','Путь к оригиналу','Доступ','Примечание','Контроль'])
choices(quotes,'G',['NET','GROSS','OTHER']);choices(quotes,'R',['PUBLIC','PRIVATE_LOCAL','RESTRICTED']);
for c in ['M','N']:choices(quotes,c,['INCLUDED','EXCLUDED','UNKNOWN'])
ql=sheet('05_Строки_КП','Исходные строки КП — без нормализации',['СтрокаКП_ID','КП_ID','Пакет_ID','Позиция / лист','Наименование / исполнение','Количество в КП','Единица КП','Цена единицы в валюте КП','Состояние цены','Включено','Исключено','Валюта (из КП)','Основа (из КП)','Источник / locator','Контроль'])
number(ql,'F');number(ql,'H');choices(ql,'I',['UNKNOWN','KNOWN','ZERO_CONFIRMED','NOT_APPLICABLE'])
kac=sheet('06_КАЦ','Сопоставление вариантов в единой валюте и составе',['Опция_ID','Группа сравнения_ID','СтрокаКП_ID','Пакет_ID','Техническая сопоставимость','Единица бюджета','Единиц КП на1 ед. бюджета','FX: база за1 валюты КП','Дата FX','Документ FX_ID','Налоговая корректировка / ед. базы','Недостающие работы / ед. базы','Доставка / ед. базы','Иные корректировки / ед. базы','Исходная цена (из строки КП)','Валюта КП','Основа КП','Сопоставимая цена / ед. базы','Выбор','Причина решения','Метод всех корректировок / единиц','Контроль','Документ корректировок_ID'])
choices(kac,'E',['UNKNOWN','COMPARABLE','NOT_COMPARABLE','DEVIATION_PENDING']);choices(kac,'S',['PENDING','SELECTED','REJECTED']);number(kac,'G',True);number(kac,'H',True)
est=sheet('07_Смета','Рабочая смета — заготовка будущей УТВ',['Строка_ID','WBS','Пакет_ID','Наименование работы / ресурса','Вид затрат','Состояние объекта','Кто поставляет','Граница затрат','Количество','Единица','Основание цены_ID','Цена / ед. базовая','Сумма базовая','Контроль готовности','Документ количества_ID','Ревизия / лист','Состояние количества','Состав / исключения','Ответственный','Решение / версия УТВ','Примечание','Тип основания цены'])
for i,q in enumerate(packages,4):
 for c,v in {1:f'EST-{i-3:03}',2:q['code_hint'],4:q['title'],17:'UNKNOWN',18:q['exclusions']}.items():est.cell(i,c,v)
choices(est,'E',['EQUIPMENT','MATERIAL','LABOUR','FABRICATION','SUBCONTRACT','DESIGN','SURVEY','TEST','LOGISTICS','FEES','RESERVE']);choices(est,'F',['EXISTING','NEW','MODIFY','REPAIR','UNKNOWN']);choices(est,'G',['OWNER','YARD','SUPPLIER','UNKNOWN']);choices(est,'H',['REMAINING','SUNK','OUTSIDE']);choices(est,'Q',['UNKNOWN','KNOWN','ZERO_CONFIRMED','NOT_APPLICABLE']);number(est,'I')
direct=sheet('06а_Оценки','Обоснованные расценки без КП — отдельный источник',['Оценка_ID','Пакет_ID','Тип основания','Описание','Единица','Цена / ед. базовая','Валюта базы','Основа NET/GROSS','Дата оценки','Документ_ID','Расчёт / метод / locator','Статус цены','Решение о применении','Контроль'])
choices(direct,'C',['INTERNAL_LABOUR','CALCULATION','HISTORICAL_COST','ALLOWANCE','FEE','RESERVE']);choices(direct,'L',['UNKNOWN','KNOWN','ZERO_CONFIRMED']);choices(direct,'H',['NET','GROSS']);number(direct,'F');choices(est,'V',['KAC','DIRECT'])
summary=sheet('08_Сводка','Сводка известной части — не утверждённая цена яхты',['WBS','Раздел','Позиций','Готово к оценке','Не закрыто','Известный остаток достройки','Известные понесённые','Известное вне бюджета','Статус'],len(raw)+5)
risks=sheet('09_Риски','Риски и обоснование резервов',['Риск_ID','WBS','Описание','Вероятность / основание','Влияние / основание','Мера снижения','Резерв базовая валюта','Метод / источник','Решение','Включён в Строка_ID','Статус','Ответственный'])
choices(risks,'K',['OPEN','MITIGATED','ACCEPTED','CLOSED']);number(risks,'G')
changes=sheet('10_Изменения','Изменения после фиксации базового бюджета',['Изменение_ID','Дата','Исходная версия УТВ','Строка_ID','Причина / решение','Прежний состав','Новый состав','Старая сумма базы','Новая сумма базы','Дельта','Влияние на срок','Статус','Утвердил / документ'])
choices(changes,'L',['PROPOSED','REVIEWED','APPROVED','REJECTED']);number(changes,'H');number(changes,'I')
control=sheet('11_Контроль','После утверждения: факт, обязательства и прогноз',['Период','WBS','Версия утверждённой УТВ','Базовый бюджет','Обязательства всего','Факт стоимости','Оплачено (справочно)','Остаток обязательств','Незаконтрактованный остаток','Утверждённый неучтённый резерв','Прогноз завершения','Отклонение от базы','Источник / решение','Контроль'])
for c in ['D','E','F','G','H','I','J']:number(control,c)
interfaces=sheet('12_Стыки','Включения и исключения — контроль двойного счёта',['Стык_ID','Пакет_ID','Операция / компонент','Статус включения','Покрывающая Строка_ID','Ответственный пакет_ID','Основание / КП / ревизия','Решение','Статус'])
choices(interfaces,'D',['INCLUDED','EXCLUDED','UNKNOWN','NOT_APPLICABLE']);choices(interfaces,'I',['OPEN','CLOSED'])
sources=sheet('13_Источники','Документы и основания — реестр для заполнения',['Документ_ID','Тип','Название','Дата','Ревизия','Путь / ссылка','SHA256','Доступ','Статус','Примечание'])
choices(sources,'H',['PUBLIC','PRIVATE_LOCAL','RESTRICTED'])
# Quote/header/line controls ensure missing inputs never become money through blank lookups.
for r in range(4,N+1):
 formula(direct,'N',r,f'=IF(A{r}="","",IF(AND(COUNTIF(A$4:A${N},A{r})=1,B{r}<>"",{count("03_Пакеты","A",f"B{r}")}=1,C{r}<>"",E{r}<>"",ISNUMBER(F{r}),F{r}>=0,G{r}<>"",G{r}=\'01_Основания\'!$B$8,H{r}<>"",H{r}=\'01_Основания\'!$B$9,I{r}<>"",J{r}<>"",{count("13_Источники","A",f"J{r}")}=1,K{r}<>"",M{r}<>"",\'01_Основания\'!$B$10<>"",OR(AND(F{r}>0,L{r}="KNOWN"),AND(F{r}=0,L{r}="ZERO_CONFIRMED"))),"READY","INCOMPLETE / BASIS"))')
 formula(quotes,'T',r,f'=IF(A{r}="","",IF(AND(COUNTIF(A$4:A${N},A{r})=1,F{r}<>"",G{r}<>"",D{r}<>"",P{r}<>"",{count("13_Источники","A",f"P{r}")}=1,OR(G{r}="NET",G{r}="GROSS",AND(G{r}="OTHER",H{r}<>""))),"READY","INCOMPLETE / ID"))')
 for col,sourcecol in [('L','F'),('M','G')]:formula(ql,col,r,f'=IF(B{r}="","",{lookup("04_КП",f"B{r}",sourcecol)})')
 formula(ql,'O',r,f'=IF(A{r}="","",IF(AND(COUNTIF(A$4:A${N},A{r})=1,{count("04_КП","A",f"B{r}")}=1,{lookup("04_КП",f"B{r}","T")}="READY",{count("03_Пакеты","A",f"C{r}")}=1,G{r}<>"",ISNUMBER(H{r}),H{r}>=0,OR(AND(H{r}>0,I{r}="KNOWN"),AND(H{r}=0,I{r}="ZERO_CONFIRMED")),N{r}<>""),"READY","INCOMPLETE / SOURCE"))')
 for col,src in [('O','H'),('P','L'),('Q','M')]:formula(kac,col,r,f'=IF(C{r}="","",IF({lookup("05_Строки_КП",f"C{r}","O")}="READY",{lookup("05_Строки_КП",f"C{r}",src)},""))',col=='O')
 ready=f'AND(COUNTIF(A$4:A${N},A{r})=1,B{r}<>"",C{r}<>"",D{r}<>"",{count("05_Строки_КП","A",f"C{r}")}=1,{lookup("05_Строки_КП",f"C{r}","C")}=D{r},{lookup("05_Строки_КП",f"C{r}","O")}="READY",E{r}="COMPARABLE",F{r}<>"",COUNT(G{r}:H{r})=2,G{r}>0,H{r}>0,I{r}<>"",J{r}<>"",{count("13_Источники","A",f"J{r}")}=1,W{r}<>"",{count("13_Источники","A",f"W{r}")}=1,COUNT(K{r}:N{r})=4,ISNUMBER(O{r}),U{r}<>"",\'01_Основания\'!$B$8<>"",OR(\'01_Основания\'!$B$9="NET",\'01_Основания\'!$B$9="GROSS"),\'01_Основания\'!$B$10<>"",IF(AND(ISNUMBER(O{r}),ISNUMBER(G{r}),ISNUMBER(H{r})),O{r}*G{r}*H{r}+SUM(K{r}:N{r})>=0,FALSE))'
 formula(kac,'V',r,f'=IF(A{r}="","",IF({ready},IF(S{r}="SELECTED",IF(AND(COUNTIFS(B$4:B${N},B{r},S$4:S${N},"SELECTED")=1,T{r}<>""),"SELECTED_READY","SELECTION_ERROR"),"COMPARABLE_READY"),"INCOMPLETE / NOT_COMPARABLE"))')
 formula(kac,'R',r,f'=IF(OR(V{r}="SELECTED_READY",V{r}="COMPARABLE_READY"),O{r}*G{r}*H{r}+SUM(K{r}:N{r}),"")',True)
 valid=f'AND(COUNTIF(A$4:A${N},A{r})=1,COUNTIF(\'02_WBS\'!$A$4:$A$200,B{r})=1,C{r}<>"",{count("03_Пакеты","A",f"C{r}")}=1,{lookup("03_Пакеты",f"C{r}","B")}=B{r},E{r}<>"",F{r}<>"",F{r}<>"UNKNOWN",G{r}<>"",G{r}<>"UNKNOWN",OR(H{r}="REMAINING",H{r}="SUNK",H{r}="OUTSIDE"),ISNUMBER(I{r}),I{r}>=0,OR(AND(I{r}>0,Q{r}="KNOWN"),AND(I{r}=0,Q{r}="ZERO_CONFIRMED")),J{r}<>"",K{r}<>"",OR(AND(V{r}="KAC",{count("06_КАЦ","A",f"K{r}")}=1,{lookup("06_КАЦ",f"K{r}","V")}="SELECTED_READY",{lookup("06_КАЦ",f"K{r}","D")}=C{r},{lookup("06_КАЦ",f"K{r}","F")}=J{r}),AND(V{r}="DIRECT",{count("06а_Оценки","A",f"K{r}")}=1,{lookup("06а_Оценки",f"K{r}","N")}="READY",{lookup("06а_Оценки",f"K{r}","B")}=C{r},{lookup("06а_Оценки",f"K{r}","E")}=J{r})),O{r}<>"",{count("13_Источники","A",f"O{r}")}=1,P{r}<>"",R{r}<>"")'
 formula(est,'N',r,f'=IF(A{r}="","",IF(Q{r}="NOT_APPLICABLE",IF(AND(COUNTIF(A$4:A${N},A{r})=1,H{r}="OUTSIDE",R{r}<>"",O{r}<>"",{count("13_Источники","A",f"O{r}")}=1),"NOT_APPLICABLE","INCOMPLETE / EXCLUSION"),IF({valid},"READY","INCOMPLETE / LINK")))')
 formula(est,'L',r,f'=IF(N{r}="READY",IF(V{r}="KAC",{lookup("06_КАЦ",f"K{r}","R")},{lookup("06а_Оценки",f"K{r}","F")}),"")',True)
 formula(est,'M',r,f'=IF(N{r}="READY",ROUND(I{r}*L{r},2),"")',True)
 formula(changes,'J',r,f'=IF(AND(A{r}<>"",COUNT(H{r}:I{r})=2),ROUND(I{r}-H{r},2),"")',True)
 formula(control,'N',r,f'=IF(A{r}="","",IF(AND(C{r}<>"",M{r}<>"",COUNT(D{r}:J{r})=7,\'01_Основания\'!$B$16<>"",\'01_Основания\'!$B$8<>""),"INPUTS_PRESENT_CHECK_OVERLAP","INCOMPLETE"))')
 formula(control,'K',r,f'=IF(N{r}="INPUTS_PRESENT_CHECK_OVERLAP",ROUND(F{r}+H{r}+I{r}+J{r},2),"")',True)
 formula(control,'L',r,f'=IF(ISNUMBER(K{r}),ROUND(K{r}-D{r},2),"")',True)
for r,q in enumerate(raw,4):
 summary.cell(r,1,q['wbs_id']);summary.cell(r,2,q['name'])
 criterion=f'A{r}&"*"'
 for col,form in [('C',f'COUNTIF(\'07_Смета\'!$B$4:$B${N},{criterion})'),('D',f'COUNTIFS(\'07_Смета\'!$B$4:$B${N},{criterion},\'07_Смета\'!$N$4:$N${N},"READY")'),('E',f'C{r}-D{r}-COUNTIFS(\'07_Смета\'!$B$4:$B${N},{criterion},\'07_Смета\'!$N$4:$N${N},"NOT_APPLICABLE")')]:formula(summary,col,r,'='+form)
 for col,bound in [('F','REMAINING'),('G','SUNK'),('H','OUTSIDE')]:
  condition=f'COUNTIFS(\'07_Смета\'!$B$4:$B${N},{criterion},\'07_Смета\'!$N$4:$N${N},"READY",\'07_Смета\'!$H$4:$H${N},"{bound}")'
  total=f'SUMIFS(\'07_Смета\'!$M$4:$M${N},\'07_Смета\'!$B$4:$B${N},{criterion},\'07_Смета\'!$H$4:$H${N},"{bound}")'
  formula(summary,col,r,f'=IF({condition}=0,"",{total})',True)
 formula(summary,'I',r,f'=IF(C{r}=0,"SCOPE_NOT_DETAILED",IF(E{r}>0,"INCOMPLETE","INPUTS_COMPLETE_SCOPE_REVIEW"))')
r=len(raw)+4;summary.cell(r,1,'ИЗВЕСТНАЯ ЧАСТЬ');summary.cell(r,2,'Не полная стоимость яхты')
for c in ['C','D','E']:formula(summary,c,r,f'=SUM({c}4:{c}{r-1})')
for c in ['F','G','H']:formula(summary,c,r,f'=IF(COUNT({c}4:{c}{r-1})=0,"",SUM({c}4:{c}{r-1}))',True)
formula(summary,'I',r,'="НЕ УТВЕРЖДЕНО — проверить состав, источники и решение владельца"')
for ws in [est,kac,ql]:
 for row in ws.iter_rows(min_row=4,max_row=N):
  for cell in row:
   if any(x in str(ws.cell(3,cell.column).value).lower() for x in ['цена','сумма','корректиров','доставка /','работы /','fx:']):cell.number_format='#,##0.00####;[Red](#,##0.00####);0.00'
# Money/quantity fields remain blank. No currencies, rates or suppliers are assumed.
a.out.parent.mkdir(parents=True,exist_ok=True);wb.save(a.out)
print(a.out)
# Schema is generated separately for reproducibility and independent checks.
schema_path=a.out.with_suffix('.schema.json');schema_path.write_text(json.dumps(schema,ensure_ascii=False,indent=2)+'\n')
