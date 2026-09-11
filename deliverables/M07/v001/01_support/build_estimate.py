"""Freeze current M03 basis into M07 working estimate; tax inline, unknowns blank."""
from pathlib import Path
from decimal import Decimal,ROUND_HALF_UP
import csv,json,hashlib
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
S=Path(__file__).resolve().parent;R=S.parents[3]
def read(path):return list(csv.DictReader((R/path).open(),delimiter=';'))
prices=read('deliverables/M03/v001/01_support/price-register.csv');P={r['candidate_id']:r for r in prices}
components=read('deliverables/M03/v001/01_support/components.csv');packages=read('deliverables/M03/v001/01_support/packages.csv');PK={r['package_id']:r for r in packages};me=read('deliverables/M01/v001/01_support/equipment-list.csv')
D=Decimal;money=lambda v:D(v).quantize(D('.01'),rounding=ROUND_HALF_UP)
def num(v):return float(v) if v not in ['',None] else None
rows=[]
for r in components:
 if not r['variant_id'].endswith('-B'):continue
 p=P[r['candidate_id']];rows.append({'id':r['line_id'],'package':r['variant_id'][:4],'me':r['me_id'],'type':'Оборудование / компонент','qty':r['quantity'],'unit':r['unit'],'name':p['manufacturer']+' '+p['model'],'source':r['candidate_id'],'net_unit':p['eur_net_unit'],'tax_mode':'SOURCE_GROSS' if r['candidate_id']=='YAN-E01' else 'PL23_ESTIMATE','rate':'' if r['candidate_id']=='YAN-E01' else '23','source_gross_eur':'58842' if r['candidate_id']=='YAN-E01' else '', 'scope_status':'OPEN','price_class':p['price_class'],'note':r['scope_note']+(' | Исходный gross включает налог неизвестной ставки; без повторного23%. NET и выделенныйVAT открыты.' if r['candidate_id']=='YAN-E01' else ' | НДС23% предварительно по поручению владельца; не установленный режим сделки.'),'url':p['price_url'],'original_price':p['listed_price'],'original_currency':p['currency']})
display_names={'EL-OPEN-MAIN':'Пусковая АКБ двигателя — тип уточняется','EL-OPEN-GEN':'Пусковые АКБ генераторов — тип уточняется','WC-RO-02':'Schenker ZEN150 — базовая цена «от», полный состав открыт'}
for r in rows:
 if r['source'] in display_names:r['name']=display_names[r['source']]
covered=set(x for r in rows for x in r['me'].split('|'))
for m in me:
 if m['item_id'] in covered:continue
 owner=next(p['package_id'] for p in packages if m['item_id'] in p['me_ids'].split('|'))
 rows.append(dict(id='OPEN-'+m['item_id'],package=owner,me=m['item_id'],type='Неоценённая функция',qty=m['quantity'],unit=m['unit'],name=m['equipment'],source='M01/'+m['item_id'],net_unit='',tax_mode='PL23_ESTIMATE',rate='23',source_gross_eur='',scope_status='CONDITIONAL' if m['item_id']=='ME-26' else 'OPEN',price_class='NO_PRICE',note=m['scope_note']+' | Резервсостава,неповтор комплектных деталей; количествоаппаратов не назначено.',url='',original_price='',original_currency=''))
gaps=[('Рабочее проектирование и интеграция МО','Чертежи фундаментов/трасс/систем; без общей архитектуры всей яхты'),('Внешние трубопроводы, кабели и монтажные материалы','Только невключённые в поставку материалы: длины, диаметры, крепления и изоляция по проекту'),('Изготовление фундаментов и локальных усилений','Внешние корпусные работы; не повтор заводских виброопор и корпусного бюджета'),('Монтаж и центровка оборудования','Труд и оснастка; без уже включённого монтажа из будущих КП'),('Настройка, пусконаладка и испытания МО','Системные/интеграционные испытания; не повтор заводских испытаний и общесудовой приёмки'),('Доставка и страхование до верфи','Маршрут, местоотгрузки,масса/габариты; есливключеноКП — обоснованныйноль'),('Пошлина и таможенное оформление','Применимость по маршруту и происхождению; не автоматическое начисление; VATбаза уточняется'),('Дополнительная установочная комплектность','Недостающие соединители,контроллеры,сервисные принадлежности послеM06; оплачиватьодинразпоBOM')]
for i,(name,note) in enumerate(gaps,1):rows.append(dict(id=f'GAP-{i:02}',package='ADD',me='',type='Работа / внешняя затрата',qty='',unit='',name=name,source='M06/scope + office/cost-policy',net_unit='',tax_mode='TAX_REVIEW' if i==7 else 'PL23_ESTIMATE',rate='' if i==7 else '23',source_gross_eur='',scope_status='OPEN',price_class='NO_PRICE',note=note,url='',original_price='',original_currency=''))
# Independent Decimal reference, preserving M03 unit EUR then rounded extended row.
for r in rows:
 net=money(D(r['qty'])*D(r['net_unit'])) if r['qty'] and r['net_unit'] else None
 vat=money(net*D(r['rate'])/100) if net is not None and r['rate'] and r['tax_mode']=='PL23_ESTIMATE' else None
 gross=money(D(r['qty'])*D(r['source_gross_eur'])) if r['qty'] and r['tax_mode']=='SOURCE_GROSS' else (net+vat if net is not None and vat is not None else None)
 r['net_eur']=str(net) if net is not None else '';r['vat_eur']=str(vat) if vat is not None else '';r['gross_eur']=str(gross) if gross is not None else ''
with (S/'estimate-lines.csv').open('w',newline='') as f:w=csv.DictWriter(f,fieldnames=list(rows[0]),delimiter=';',lineterminator='\n');w.writeheader();w.writerows(rows)
stats={'main_rows':len(rows),'source_B_rows':25,'missing_ME_added':5,'additional_unpriced_cost_rows':8,'priced_net_rows':sum(bool(r['net_eur']) for r in rows),'priced_gross_rows':sum(bool(r['gross_eur']) for r in rows),'unpriced_gross_rows':sum(not bool(r['gross_eur']) for r in rows),'known_net_eur':str(sum((D(r['net_eur']) for r in rows if r['net_eur']),D(0))),'estimated_vat_identified_part_eur':str(sum((D(r['vat_eur']) for r in rows if r['vat_eur']),D(0))),'source_gross_tax_not_split_eur':'58842.00','known_gross_mixed_basis_eur':str(sum((D(r['gross_eur']) for r in rows if r['gross_eur']),D(0))),'status':'FROZEN_WORKING_ESTIMATE_PARTIAL_SCOPE_NOT_UTV'}
(S/'totals.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2)+'\n')
wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
def sheet(name,heads,data,widths):
 ws=wb.create_sheet(name);ws.append(heads)
 for row in data:ws.append(row)
 ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.zoomScale=80
 for row in ws:
  ws.row_dimensions[row[0].row].height=65 if row[0].row>1 else 45
  for c in row:
   c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=11)
   if c.row==1:c.fill=PatternFill('solid',fgColor='234351');c.font=Font(name='Calibri',color='FFFFFF',bold=True)
   elif c.data_type=='f':c.fill=PatternFill('solid',fgColor='E5EBEF')
   elif c.row%2==0:c.fill=PatternFill('solid',fgColor='F1F5F7')
   if isinstance(c.value,(float,int)) or c.data_type=='f':c.number_format='#,##0.00'
   if isinstance(c.value,str) and c.value.startswith('https://'):c.hyperlink=c.value;c.font=Font(color='006DA3',underline='single')
 for i in range(1,len(heads)+1):ws.column_dimensions[get_column_letter(i)].width=widths.get(i,22)
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1';return ws
main=[]
for i,r in enumerate(rows,2):
 j=f'=IF(AND(ISNUMBER(E{i}),ISNUMBER(I{i})),ROUND(E{i}*I{i},2),"")'
 m=f'=IF(AND(K{i}="PL23_ESTIMATE",ISNUMBER(J{i}),ISNUMBER(L{i})),ROUND(J{i}*L{i}/100,2),"")'
 o=f'=IF(K{i}="SOURCE_GROSS",IF(AND(ISNUMBER(E{i}),ISNUMBER(N{i})),ROUND(E{i}*N{i},2),""),IF(AND(ISNUMBER(J{i}),ISNUMBER(M{i})),ROUND(J{i}+M{i},2),""))'
 main.append([r['id'],r['package'],r['me'],r['type'],num(r['qty']),r['unit'],r['name'],r['source'],num(r['net_unit']),j,r['tax_mode'],num(r['rate']),m,num(r['source_gross_eur']),o,r['scope_status'],r['price_class'],r['note'],r['url'],num(r['original_price']),r['original_currency']])
ws=sheet('01_Смета',['Строка','Пакет','Группы M01','Вид затрат','Количество','Ед.','Наименование / исполнение','Источник ID','Цена NET EUR/ед.','Стоимость NET EUR','Основание НДС','НДС % предварит.','Сумма НДС EUR','Исходный gross EUR/ед.','Стоимость с НДС EUR','Комплектность','Статус цены','Примечание / исключения','Прямая ссылка','Цена в источнике','Валюта источника'],main,{3:24,7:50,8:24,11:26,18:100,19:45})
for col in ['E','I','K','L','N']:
 for c in ws[col][1:]:c.font=Font(color='0066AA')
last=ws.max_row;tot=last+1
ws.cell(tot,7,'Известная оценённая часть, НЕ полный бюджет')
for col in ['J','M','O']:ws[f'{col}{tot}']=f'=SUM({col}2:{col}{last})';ws[f'{col}{tot}'].number_format='#,##0.00'
ws.cell(tot,18,'NET и выделенныйVAT не включаютYanmarSOURCE_GROSS; суммаgrossсмешиваетPL23сценарий иценупродавцаснеизвестнойставкой. Остальныепустыезатраты сверхсуммы.')
ws.row_dimensions[tot].height=75
summary=[]
for p in packages+[dict(package_id='ADD',title='Монтаж и внешние затраты')]:
 pk=p['package_id']; rr=[i for i,r in enumerate(rows,2) if r['package']==pk];e=lambda col:'=IF(COUNT('+','.join(f"'01_Смета'!{col}{i}" for i in rr)+')=0,"",SUM('+','.join(f"'01_Смета'!{col}{i}" for i in rr)+'))'
 summary.append([pk,p['title'],e('J'),e('M'),e('O'),len(rr)-sum(bool(rows[i-2]['gross_eur']) for i in rr),'НЕПОЛНЫЙ','Цена оборудования и компонента не закрывает установочный пакет'])
summ=sheet('00_Итог',['Пакет','Система','Известная NET EUR','НДС выделен / расчёт EUR','Оценено с НДС EUR','Строк без суммы','Статус','Примечание'],summary,{2:38,3:24,4:28,5:27,6:20,7:22,8:75})
z=summ.max_row+1;summ.cell(z,2,'ОЦЕНЁННАЯ ЧАСТЬ — НЕ ПОЛНАЯ СТОИМОСТЬ')
for col in ['C','D','E']:summ[f'{col}{z}']=f'=SUM({col}2:{col}{z-1})';summ[f'{col}{z}'].number_format='#,##0.00'
summ.cell(z,6,stats['unpriced_gross_rows']);summ.row_dimensions[z].height=60
notes=[('Статус','M07/v001 — выпуск текущей рабочей сметы по имеющимся сведениям. Не утверждённая УТВ, не заказ, не полный бюджет установленного МО.'),('НДС в строках','По уточнению владельца23% включены предварительно для известныхNET; SOURCE_GROSS уже сналогом, без повторного начисления. Налоговый вычет не применён.'),('Смешанный gross','Yanmar58842EUR включает налог неизвестнойставки. Его NET/VATневыделены. Общая оценённаяgrossсумма не единая унифицированная польская налоговая база и не график платежей.'),('Открытые затраты','Нулевых цен для неизвестного нет. Пустые количества/стоимости требуютданных; полныйитог незадан. Монтаж/материалы/трассы/таможня/ПНР не закрыты.'),('Варианты','Рабочая ветвьB сохранена изM03: Yanmar/Panda. Это не выбор закупки и не средний ценовойранг. A/C показаны в отдельном сравнении; альтернативы не складываются.'),('Размещение','По поручению владельца дальнейшаякомпоновка по сечениям проекта отдельно; G03не блокирует фиксациюэтогосостояниясметы и не подтверждает fit.'),('Предыдущие версии','U01/M01–M06/G03/EXP003 сохранены; сметные числа не перенесены в утверждённыйбюджет.'),('Резерв','Процент риска, накладных и установкиневыдуман. Неоценённыеработы не компенсированы произвольнойнадбавкой.')]
for key,val in notes:z+=1;summ.cell(z,2,key);summ.cell(z,3,val);summ.merge_cells(start_row=z,start_column=3,end_row=z,end_column=8);summ.cell(z,3).alignment=Alignment(wrap_text=True,vertical='top');summ.row_dimensions[z].height=50
wb._sheets.remove(summ);wb._sheets.insert(0,summ)
# Freeze all three alternatives as comparison only; no tax arithmetic or synthetic rank.
alt=[]
for p in packages:
 row=[p['package_id'],p['title']]
 for letter in 'ABC':
  cr=[r for r in components if r['variant_id']==p['package_id']+'-'+letter];vals=[money(D(r['quantity'])*D(P[r['candidate_id']]['eur_net_unit'])) for r in cr if r['quantity'] and P[r['candidate_id']]['eur_net_unit']]
  row += [float(sum(vals,D(0))) if vals else None,len(cr)-len(vals)]
 row+=['Только известная NETчасть; состав/налоги различны; неранжировать и нескладывать альтернативы'];alt.append(row)
sheet('02_Сравнение',['Пакет','Система','A известная NET','A без оценки','B известная NET','B без оценки','C известная NET','C без оценки','Условие'],alt,{2:40,9:95})
sheet('03_Источники',list(prices[0]),[[num(r[k]) if k in ['net_unit_price','eur_net_unit','listed_price','qty'] else r[k] for k in prices[0]] for r in prices],{5:48,10:42,13:50,14:50,20:100})
subs=read('deliverables/M05/v001/01_support/price-substitutions.csv');sheet('04_Подстановки_M05',list(subs[0])+['Включение в основную смету'],[[r[k] for k in subs[0]]+['НЕ ДОБАВЛЕНО: альтернативный6LPA; не повторять комплектный8LV'] for r in subs],{2:40,16:100,19:80})
scope=read('deliverables/M06/v001/01_support/scope-register.csv');sheet('05_Комплектность_M06',list(scope[0]),[[r[k] for k in scope[0]] for r in scope],{3:48,10:60,13:75,14:70})
basis=notes+[('Первичная ставкаPL','https://www.podatki.gov.pl/podatki-firmowe/vat/stawki-i-limity'),('Основание источников','M03/v001 цены11.09.2026 и справочныйFX10.09.2026. M05подстановки отдельно. Новый поиск цен не выполнялся.'),('Округление','ИсходныйNET→EURценаединицы какM03; количество×цена округляетсяHALF_UP доцента, затемVATстроки доцента, итогсуммируетстроки.'),('Налоги и маршрут','PL23 — плановое допущение владельца. WNT/импорт/вычет и счёт продавца потребуют проверки; импортнаяналоговаябаза включает свои элементы, простойтоварныйNETеёнеопределяет.'),('Состав команды','Propulsion:experiment_basis;охват:tools_freecad;сведение:root;независимыйвыпуск:experiment_model')]
sheet('06_Основания',['Тема','Основание / предел'],basis,{1:30,2:150})
wb.save(S.parent/'00_engine-room-estimate.xlsx');print(json.dumps(stats,ensure_ascii=False))
