"""Build a working KAC; source releases remain immutable.

Requires openpyxl. Monetary source units use HALF_UP cents, then EUR cents;
component extensions multiply the displayed EUR unit. Recalculate a disposable
copy with LibreOffice, inspect it, then release that calculated XLSX.
"""
import csv, json, re
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

S=Path(__file__).resolve().parent
R=S.parents[3]
M02=R/'deliverables/M02/v001/01_support'
D=Decimal
FX=D(json.loads((M02/'fx.json').read_text())['quote_per_base'])
def cash(x): return D(str(x)).quantize(D('.01'),rounding=ROUND_HALF_UP)
def read(p):
    with p.open(newline='',encoding='utf-8-sig') as f:return list(csv.DictReader(f,delimiter=';'))
def write(p,rows):
    with p.open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]),delimiter=';',lineterminator='\n');w.writeheader();w.writerows(rows)
def n(x):return float(x) if x not in ['',None] else None

prices=[]
for area,path in [('M02',M02/'all-candidates.csv'),('Yanmar',S/'yanmar/candidates.csv'),('Panda',S/'panda/candidates.csv')]:
    for old in read(path):
        r={k:old.get(k,'') for k in read(M02/'propulsion/candidates.csv')[0]}
        r['me_id']=re.sub(r'ME-?(\d{2})',r'ME-\1',r['me_id'])
        state='NO_VERIFIED_NET'
        if r['net_unit_price']:
            state='OBSERVED_NET_CONDITIONAL_SCOPE'
            if 'HISTORICAL' in r['technical_status']:state='HISTORICAL_NOT_CURRENT'
            elif 'FROM' in r['tax_basis']:state='FROM_INCOMPLETE_CONFIGURATION'
            elif 'PRICE_VARIATION' in r['technical_status']:state='VARIABLE_PRICE_RECONFIRM'
        r['price_class']=state
        r['evidence_type']='PUBLIC_PRICE_NOT_KP'
        r['price_date']='2023-03' if state=='HISTORICAL_NOT_CURRENT' else 'UNKNOWN'
        r['source_revision']=str(path.relative_to(R))
        r['eur_net_unit']=''
        if r['net_unit_price'] and r['currency'] in ['EUR','GBP']:
            val=cash(r['net_unit_price']);r['eur_net_unit']=str(cash(val/FX) if r['currency']=='GBP' else val)
        prices.append(r)
byid={r['candidate_id']:r for r in prices}
assert len(prices)==len(byid)
write(S/'price-register.csv',prices)
packages=read(S/'packages.csv')
assert [p['package_id'] for p in packages]==[f'PK{i:02}' for i in range(1,13)]
me=[m for p in packages for m in p['me_ids'].split('|')]
assert len(me)==29 and len(set(me))==29

yan=next(r['candidate_id'] for r in prices if r['candidate_id'].startswith('YAN-') and '8LV320' in r['model'])
panda=next((r['candidate_id'] for r in prices if r['candidate_id'].startswith('PAN-') and '12000' in r['model'] and r['currency']=='EUR' and r['net_unit_price']), 'PROP-G01')
# These are starting design hypotheses. A/B/C are not price or quality tiers.
choices={
 'PK01':[['PROP-E02'],[yan],['PROP-E04']],
 'PK02':[['PROP-G03'],[panda],['PROP-G02']],
 'PK03':[['EL-INV-A','EL-ISO-A'],['EL-INV-B','EL-ISO-B'],['EL-INV-C','EL-ISO-B']],
 'PK04':[['EL-BAT-A','EL-OPEN-MAIN','EL-OPEN-GEN','EL-REM-A'],['EL-BAT-B','EL-OPEN-MAIN','EL-OPEN-GEN','EL-REM-A'],['EL-BAT-C','EL-OPEN-MAIN','EL-OPEN-GEN','EL-REM-A']],
 'PK05':[['WC-DC-01','WC-AC-01'],['WC-DC-02','WC-AC-02'],['WC-DC-03','WC-AC-03']],
 'PK06':[['WC-RO-01'],['WC-RO-02'],['WC-RO-03']],
 'PK07':[['WC-HW-01'],['WC-HW-02'],['WC-HW-03']],
 'PK08':[['WC-CH-01','WC-AUX-01'],['WC-CH-02','WC-AUX-01'],['WC-CH-03','WC-AUX-01']],
 'PK09':[['A09','A10','A15','A16'],['A09','A11','A15','A16'],['A09','A11','A15','A16']],
 'PK10':[['A12','A13','A14','A20']]*3,
 'PK11':[['A17','A19'],['A18','A19'],['A18','A19']],
 'PK12':[['A01','A02','A03','A04','A05'],['A07'],['A08']]}
variants=[];components=[];checks=[]
for p in packages:
    for j,letter in enumerate('ABC'):
        vid=p['package_id']+'-'+letter;ids=choices[p['package_id']][j]
        variants.append({'variant_id':vid,'package_id':p['package_id'],'description':' + '.join(byid[x]['model'] for x in ids),
          'fixed_quantities':p['fixed_quantities'],'technical_status':'NOT_CHECKED','commercial_status':'PUBLIC_ONLY',
          'tax_completeness':'OPEN','selection':'NOT_SELECTED','open_inputs':p['open_inputs']})
        for k,cid in enumerate(ids,1):
            r=byid[cid]
            components.append({'line_id':vid+f'-L{k:02}','variant_id':vid,'me_id':r['me_id'],'candidate_id':cid,
              'quantity':r['qty'],'unit':r['unit'],'scope_note':r['scope_notes']})
        # No physical quantity is invented for missing scope: this is a checklist.
        for k,scope in enumerate(p['scope_to_complete'].split('|'),1):
            checks.append({'check_id':vid+f'-S{k:02}','variant_id':vid,'required_scope':scope.strip(),
              'status':'OPEN','closure_evidence':'','owner':'','cost_location':''})
        checks.append({'check_id':vid+'-EXT','variant_id':vid,
          'required_scope':'Подтвердить полный построчный BOM и границы: нет отсутствующих платных частей или внешних работ; у каждого включения/исключения есть место стоимости. Все монтажные доплаты учтены один раз.',
          'status':'OPEN','closure_evidence':'','owner':'','cost_location':''})
write(S/'variant-register.csv',variants);write(S/'components.csv',components);write(S/'scope-checklist.csv',checks)

wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
def sheet(name,headers,rows,widths=None):
    ws=wb.create_sheet(name);ws.append(headers)
    for row in rows:ws.append(row)
    ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.zoomScale=75
    ws.row_dimensions[1].height=45
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor='17384D');c.font=Font(color='FFFFFF',bold=True)
        c.alignment=Alignment(wrap_text=True,vertical='center')
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height=66
        for c in row:
            c.alignment=Alignment(wrap_text=True,vertical='top')
            if c.data_type=='f':c.fill=PatternFill('solid',fgColor='E8ECEF')
            elif c.row%2==0:c.fill=PatternFill('solid',fgColor='F1F6FA')
            if isinstance(c.value,(int,float)):c.number_format='#,##0.00'
            if isinstance(c.value,str) and c.value.startswith('https://'):
                c.hyperlink=c.value;c.font=Font(color='006CAA',underline='single')
    for i in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(i)].width=(widths or {}).get(i,23)
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
    return ws
def inputs(ws,cols):
    for col in cols:
        for c in ws[col][1:]:c.font=Font(color='0066BB');c.fill=PatternFill('solid',fgColor='EFF8FF')
def enum(ws,col,options):
    d=DataValidation(type='list',formula1='"'+','.join(options)+'"',allow_blank=True);d.errorTitle='Значение из списка';d.error='Выберите допустимый статус';d.showErrorMessage=True
    ws.add_data_validation(d);d.add(f'{col}2:{col}{ws.max_row}')
def decimal(ws,col,upper=None):
    d=DataValidation(type='decimal',operator='between' if upper is not None else 'greaterThanOrEqual',formula1=0,formula2=upper,allow_blank=True)
    d.showErrorMessage=True;d.error='Нужно неотрицательное число в допустимом диапазоне';ws.add_data_validation(d);d.add(f'{col}2:{col}{ws.max_row}')

intro=[
 ['Статус','M03 / 11.09.2026 — рабочая КАЦ с открытыми исходными данными. 36 слотов A/B/C для 12 функциональных пакетов; это не 36 полученных КП и не три утверждённых комплекта яхты.'],
 ['Yanmar и Fischer Panda','Yanmar добавлен отдельными новыми источниками. Fischer Panda сохранён, добавлены предложения/исполнения. Наличие марки не означает выбор или совместимость с сетью/пуском.'],
 ['Сравниваем результат','Разные устройства и внутренний BOM допустимы при нужном результате. Основные заданные количества сохраняются. A/B/C — гипотезы исполнения, не ранг цены или качества.'],
 ['Порядок работы','Согласовать результат и BOM → дополнить источники цен/КП → закрыть комплектность со ссылками → проверить технику/коммерческие условия → доставка и налоги → решение. Слоты PK09–PK11 пока используют общие ориентиры, не три независимые поставки.'],
 ['Листы','01: функциональное задание; 02: цены и прямые URL; 03: построчный состав вариантов; 04: полнота; 05: КАЦ; 06: налоговые операции; 07: открытые вопросы.'],
 ['Известная часть','На листе 05 сумма известных товаров дана отдельно. Полный NET появится только после закрытия состава, цен, исключений источников и двух допусков. Пустая ячейка не означает ноль. Общего итога МО нет.'],
 ['Количество','1 двигатель; 2 генератора; 2 основных инвертора; 6+2+1 физических АКБ; 1 водяная установка каждого питания; 1 опреснитель; 2 чиллера. ME06 — 1 функция резерва, не обязательный третий инвертор. Количество бойлеров и обвязки открыто.'],
 ['Цена','NET в исходной валюте → единичный EUR с округлением HALF_UP до цента → количество. GBP / 0.85915, ECB 10.09.2026. Курс для сравнения, не автоматический курс налоговой декларации.'],
 ['Доплаты','В 05_КАЦ шесть полей доплат: доставка, страхование, пошлина, оформление, монтаж/ПНР и прочее. Ввести обоснованный ноль, если затрата отсутствует или уже включена. Не учитывать повторно включённую услугу.'],
 ['НДС','Покупатель, VAT-UE и использование яхты ещё неизвестны. На 06_НДС нет ставки 23% по умолчанию и автоматического вычета. Операции разделять по счетам/таможенным документам; не применять одну ставку ко всему варианту без основания.'],
 ['НДС и деньги','Невозмещаемый VAT входит в стоимость, вычет показывается отдельно. Денежный платёж и ожидаемый возврат вводятся отдельно: при WNT зачёт может быть без оплаты VAT. График финансирования ещё не рассчитан.'],
 ['Статусы закрытия','PASS/CONFIRMED_FULL/CLOSED/CONFIRMED_ALL_DOCUMENTS требуют реального подтверждения, не проставляются по наличию числа. Внешние корпусные и общесудовые работы раскрывать отдельными связями.'],
 ['Редактирование','Синие поля — ввод, серые — формулы. Регистр цен содержит источники, а не выбранную закупку. После выпуска работать в следующей версии/копии; расширяя листы, продлить формулы и диапазоны. Формулы в этом выпуске рассчитаны LibreOffice.'],
 ['Сохранность','Исходные U01/M01/M02 и чертежи сохранены. Полные сторонние PDF локально .local/M03, ссылки и SHA в сопровождении. Поставщикам ничего не отправлено.']]
sheet('00_Старт',['Тема','Правило'],intro,{1:28,2:145})
sheet('01_Пакеты',['Пакет','Результат / система','Группы M01','Требуемый выход','Фиксированные количества','Полная граница поставки','Критерий приёмки','Открытые параметры'],
 [[p[k] for k in packages[0]] for p in packages],{2:30,3:30,4:70,5:45,6:120,7:95,8:100})
source_rows=[]
for rownum,r in enumerate(prices,2):
    dims=' × '.join(r[k] for k in ['length_mm','width_mm','height_mm']) if all(r[k] for k in ['length_mm','width_mm','height_mm']) else ''
    eur_formula=f'=IF(AND(ISNUMBER(K{rownum}),K{rownum}>=0,G{rownum}="EUR"),ROUND(K{rownum},2),IF(AND(ISNUMBER(K{rownum}),K{rownum}>=0,G{rownum}="GBP"),ROUND(ROUND(K{rownum},2)/0.85915,2),""))'
    source_rows.append([r['candidate_id'],r['me_id'],r['manufacturer'],r['model'],n(r['qty']),r['unit'],r['currency'],n(r['listed_price']),r['tax_basis'],n(r['tax_rate']),n(r['net_unit_price']),eur_formula,r['price_class'],r['evidence_type'],r['price_date'],r['access_date'],r['price_url'],r['spec_url'],dims,n(r['mass_kg']),r['technical_status'],r['scope_notes'],r['availability'],r['source_revision']])
sheet('02_Цены',['Источник_ID','M01','Производитель','Исполнение','Кол-во ориентир','Ед.','Валюта','Исходная цена','Налоговая основа','Ставка источника %','NET / ед. валюта','NET / ед. EUR','Класс цены','Тип доказательства','Дата цены','Дата просмотра','Прямая цена','Спецификация','Габарит изделия мм','Масса кг','Технический статус','Включения и ограничения','Наличие','Происхождение записи'],source_rows,{4:48,9:38,13:40,14:33,17:40,18:40,21:48,22:110,23:42,24:55})

cp=[]
for i,r in enumerate(components,2):
    # INDEX on an empty source cell may become zero: explicitly retain blank.
    lookup=lambda col:f'IFERROR(IF(INDEX(\'02_Цены\'!${col}$2:${col}$1000,MATCH(D{i},\'02_Цены\'!$A$2:$A$1000,0))="","",INDEX(\'02_Цены\'!${col}$2:${col}$1000,MATCH(D{i},\'02_Цены\'!$A$2:$A$1000,0))),"")'
    cp.append([r['line_id'],r['variant_id'],r['me_id'],r['candidate_id'],n(r['quantity']),r['unit'],
      '='+lookup('D'),'='+lookup('L'),f'=IF(AND(ISNUMBER(E{i}),E{i}>=0,ISNUMBER(H{i})),ROUND(E{i}*H{i},2),"")',
      '='+lookup('M'),f'=IF(AND(ISNUMBER(I{i}),J{i}="OBSERVED_NET_CONDITIONAL_SCOPE",OR({lookup("E")}="",E{i}={lookup("E")})),"PRICE_OBSERVED","BLOCK")',r['scope_note']])
ws=sheet('03_Состав',['Строка','Вариант','M01 / интерфейсы','Источник_ID','Количество','Ед.','Исполнение из источника','NET / ед. EUR','Известная NET строки','Класс цены','Допуск ценового основания','Состав / недостающее'],cp,{3:30,4:25,7:48,10:40,11:30,12:110})
inputs(ws,['D','E','F','L']);decimal(ws,'E')
scope_rows=[[r[k] for k in checks[0]] for r in checks]
ws=sheet('04_Комплектность',['Проверка_ID','Вариант','Что должно быть закрыто','Статус','Документ / лист закрытия','Ответственный','Где учтена стоимость'],scope_rows,{3:130,5:50,7:60})
inputs(ws,['D','E','F','G']);enum(ws,'D',['OPEN','CLOSED'])

tax_rows=[]
for i,v in enumerate(variants,2):
    tax_rows.append(['TAX-'+v['variant_id'],v['variant_id'],None,None,None,None,None,None,None,
      f'=IF(AND(ISNUMBER(H{i}),H{i}>=0,ISNUMBER(I{i}),I{i}>=0,I{i}<=100),ROUND(H{i}*I{i}/100,2),"")',None,
      f'=IF(AND(ISNUMBER(J{i}),ISNUMBER(K{i}),K{i}>=0,K{i}<=100),ROUND(J{i}*K{i}/100,2),"")',
      f'=IF(AND(ISNUMBER(J{i}),ISNUMBER(L{i})),ROUND(J{i}-L{i},2),"")',None,None,None,None,None,
      f'=IF(AND(C{i}<>"",D{i}<>"",E{i}<>"",F{i}<>"",G{i}<>"",R{i}<>"",ISNUMBER(M{i})),"CALCULATED_WITH_INPUTS","OPEN")'])
ws=sheet('06_НДС',['Налоговая строка','Вариант','Счёт / операция / компоненты','Покупатель / VAT статус','Фактическая отправка','Страна VAT / импортёр','Режим','Налоговая база EUR','Ставка %','Начисленный VAT EUR','Подтверждённая доля вычета %','VAT к вычету EUR','Невозмещаемый VAT EUR','VAT: денежный платёж EUR','Ожидаемый денежный возврат EUR','Дата оплаты','Дата возврата / зачёта','Правовое и документальное основание','Контроль расчёта'],tax_rows,{3:55,4:45,5:35,6:40,7:30,18:75,19:34})
inputs(ws,['C','D','E','F','G','H','I','K','N','O','P','Q','R'])
for col in ['H','N','O']:decimal(ws,col)
for col in ['I','K']:decimal(ws,col,100)
enum(ws,'G',['PL_DOMESTIC','WNT','EU_B2C','IMPORT_PL','OTHER_DOCUMENTED'])

cv=[]
for i,v in enumerate(variants,2):
    vid=v['variant_id'];end=len(components)+1;ce=len(checks)+1
    c=f"'03_Состав'!$B$2:$B${end}";vals=f"'03_Состав'!$I$2:$I${end}"
    sc=f"'04_Комплектность'!$B$2:$B${ce}";scst=f"'04_Комплектность'!$D$2:$D${ce}"
    scdoc=f"'04_Комплектность'!$E$2:$E${ce}";sccost=f"'04_Комплектность'!$G$2:$G${ce}"
    # Zero is a known value only after source/quantity evidence, never a blank.
    cv.append([vid,v['package_id'],v['description'],v['fixed_quantities'],
      f'=COUNTIF({c},A{i})',f'=SUMPRODUCT(({c}=A{i})*ISNUMBER({vals}))',
      f'=IF(F{i}>0,ROUND(SUMIF({c},A{i},{vals}),2),"")',
      f'=SUMPRODUCT(({sc}=A{i})*((({scst}<>"CLOSED")+({scdoc}="")+({sccost}=""))>0))',
      f'=COUNTIFS({c},A{i},\'03_Состав\'!$K$2:$K${end},"BLOCK")',
      v['technical_status'],v['commercial_status'],
      f'=IF(AND(E{i}>0,E{i}=F{i},H{i}=0,I{i}=0,J{i}="PASS",K{i}="CONFIRMED_FULL"),G{i},"")',
      None,None,None,None,None,None,
      f'=IF(AND(ISNUMBER(L{i}),COUNT(M{i}:R{i})=6),ROUND(SUM(L{i}:R{i}),2),"")',
      f'=IF(AND(V{i}="CONFIRMED_ALL_DOCUMENTS",COUNTIF(\'06_НДС\'!$B$2:$B$1000,A{i})>0,COUNTIFS(\'06_НДС\'!$B$2:$B$1000,A{i},\'06_НДС\'!$S$2:$S$1000,"<>CALCULATED_WITH_INPUTS")=0),ROUND(SUMIF(\'06_НДС\'!$B$2:$B$1000,A{i},\'06_НДС\'!$M$2:$M$1000),2),"")',
      f'=IF(AND(ISNUMBER(S{i}),ISNUMBER(T{i})),ROUND(S{i}+T{i},2),"")',v['tax_completeness'],v['selection'],v['open_inputs'],
      f'=IF(ISNUMBER(U{i}),"CALCULATED_FOR_REVIEW","DRAFT_OPEN")'])
ws=sheet('05_КАЦ',['Вариант','Пакет','Начальная гипотеза исполнения','Заданные количества','Строк состава','Строк с суммой','Известная часть NET EUR','Открыто по составу','Недопущенных цен','Технический допуск','Коммерческий состав','Полный NET состава EUR','Доставка NET EUR','Страхование NET EUR','Пошлина EUR','Оформление NET EUR','Монтаж / ПНР NET EUR','Прочие NET EUR','Полная стоимость NET EUR','Невозмещаемый VAT EUR','Стоимость проекта EUR','Полнота налоговых документов','Выбор владельца','Что ещё требуется','Статус'],cv,{3:70,4:48,10:26,11:32,22:38,24:110,25:32})
inputs(ws,['J','K','M','N','O','P','Q','R','V','W','X'])
enum(ws,'J',['NOT_CHECKED','PASS','FAIL']);enum(ws,'K',['PUBLIC_ONLY','INCOMPLETE','CONFIRMED_FULL'])
enum(ws,'V',['OPEN','CONFIRMED_ALL_DOCUMENTS']);enum(ws,'W',['NOT_SELECTED','PREFERRED_FOR_REVIEW','OWNER_APPROVED'])
for col in ['M','N','O','P','Q','R']:decimal(ws,col)

sheet('07_Открытые вопросы',['Пакет','Результат','Требуется для допуска','Первое действие'],
 [[p['package_id'],p['title'],p['open_inputs'],'Согласовать требуемый выход, проверить обслуживаемую компоновку и полный BOM; запросить недостающие основания.'] for p in packages],{2:35,3:130,4:80})
wb._sheets.sort(key=lambda w:w.title)
for ws in wb:
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.data_type=='f':c.number_format='#,##0.00'
wb.save(S.parent/'00_working-KAC.xlsx')
(S/'build-stats.json').write_text(json.dumps({'sources':len(prices),'packages':12,'variants':36,'component_rows':len(components),'scope_rows':len(checks),'worksheets':8,'yanmar_primary':yan,'panda_primary':panda,'status':'DRAFT_OPEN_NO_COMPLETE_TOTAL'},indent=2)+'\n')
print((S/'build-stats.json').read_text())
