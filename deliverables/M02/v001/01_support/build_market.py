"""Build the M02 research workbook from reviewed registers; no supplier requests.

Run with openpyxl installed. Does not change source registers or prior releases.
Money policy: round source NET/unit to cents; convert that unit to EUR and round
HALF_UP; multiply displayed EUR/unit by the fixed quantity. No basket grand total.
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv, json, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

S = Path(__file__).resolve().parent
R = S.parents[3]
D = Decimal
FX = D(json.loads((S/'fx.json').read_text())['quote_per_base'])
def money(x): return D(str(x)).quantize(D('.01'), rounding=ROUND_HALF_UP)
def read(p):
    with p.open(newline='', encoding='utf-8-sig') as f: return list(csv.DictReader(f, delimiter=';'))
def write(p, rows):
    with p.open('w', newline='', encoding='utf-8') as f:
        w=csv.DictWriter(f, fieldnames=list(rows[0]), delimiter=';', lineterminator='\n'); w.writeheader(); w.writerows(rows)

rows=[]
for area in ['propulsion','electrical','water-climate','auxiliary']:
    for original in read(S/area/'candidates.csv'):
        r=dict(original)
        r['me_id']=re.sub(r'ME-?(\d{2})', r'ME-\1', r['me_id'])
        r['source_register']=f'{area}/candidates.csv'
        r['commercial_type']='PUBLIC_PRICE_NOT_SUPPLIER_QUOTE'
        r['price_date']='2023-03' if 'HISTORICAL' in r['technical_status'] else 'UNKNOWN; accessed 2026-09-11'
        r['price_class']='NO_VERIFIED_NET'
        if r['net_unit_price']:
            r['price_class']='OBSERVED_NET_CONDITIONAL_SCOPE'
            if 'HISTORICAL' in r['technical_status']: r['price_class']='HISTORICAL_2023_NOT_CURRENT'
            elif 'FROM' in r['tax_basis']: r['price_class']='FROM_INCOMPLETE_CONFIGURATION'
            elif 'PRICE_VARIATION' in r['technical_status']: r['price_class']='VARIABLE_PRICE_RECONFIRM'
        r['eur_net_unit']=''; r['eur_fixed_qty_line']=''
        if r['net_unit_price'] and r['currency'] in ['EUR','GBP']:
            n=money(r['net_unit_price'])
            eur=money(n/FX) if r['currency']=='GBP' else n
            r['eur_net_unit']=str(eur)
            if r['qty']: r['eur_fixed_qty_line']=str(money(eur*D(r['qty'])))
        r['scope_qualification']='TECHNICAL_OUTPUT_AND_FULL_INSTALLED_SCOPE_OPEN'
        if r['candidate_id'].startswith('PROP-E') and 'ME-03' in r['me_id']:
            r['scope_qualification']+='; ME-03 PARTIAL'
        if r['candidate_id'].startswith('PROP-G'):
            r['scope_qualification']+='; ME-19 PARTIAL'
        rows.append(r)
assert len(rows)==65 and len({r['candidate_id'] for r in rows})==65
byid={r['candidate_id']:r for r in rows}
write(S/'all-candidates.csv', rows)

# Known quantities are not enlarged to fill price gaps.
fixed={'ME-01':'1','ME-02':'1','ME-04':'2','ME-05':'2','ME-06':'1','ME-07':'1','ME-08':'6','ME-11':'1','ME-12':'1','ME-13':'1','ME-15':'2','ME-27':'2'}
for r in rows:
    primary=r['me_id'].split('|')[0]
    if primary in fixed: assert r['qty']==fixed[primary], r['candidate_id']
    if r['net_unit_price'] and r['listed_price'] and r['tax_rate']:
        expected=money(D(r['listed_price'])/(1+D(r['tax_rate'])/100))
        assert abs(expected-money(r['net_unit_price']))<=D('.01'), (r['candidate_id'], expected)
    if not r['net_unit_price']: assert not r['eur_net_unit'] and not r['eur_fixed_qty_line']

gaps={
1:'Нет трёх текущих NET. Нужны сопротивление/скорость, рабочий рейтинг, PTO, 24 В и полный комплект.',
2:'Цена частично в моторном пакете. Отдельный редуктор повторно не добавлять; отношение и PTO открыты.',
3:'Полного валопровода нет: вал, дейдвуд, уплотнение, муфта, упор, монтаж и согласование винта.',
4:'Три наблюдаемые NET есть; полезная мощность, резерв при отказе, пусковое напряжение и обвязка не уравнены.',
5:'Три цены условного класса. Нужен баланс нагрузок, пуск, заряд, разделение ветвей и допустимый параллельный режим.',
6:'Одна резервная функция, не автоматически третий инвертор. Мощность насоса и схема открыты.',
7:'Одна АКБ; напряжение/CCA/химия и выбранный двигатель открыты, цены нет.',
8:'Три цены только условных 115 А·ч / 12 В. Владелец не подтвердил единицу и напряжение; химия/автономность открыты.',
9:'Две NET и одна цена с неизвестной ставкой. Ток берега, число устройств и электрическая схема открыты.',
10:'Цены отдельных разных компонентов, не трёх щитов. Полные 24 В DC / AC щиты, защиты, заряд и кабели открыты.',
11:'Две NET, третий RFQ. Рабочая точка и состав, точное соответствие OEM/SKU открыты.',
12:'Одна NET, два RFQ. Напор/расход в точках потребления и комплектность открыты.',
13:'Базовая NET, цена ОТ, третий RFQ. Довести каждую систему до нужной воды, промывки и контроля качества.',
14:'Три единичные NET для примеров 40 л. Количество и потребность горячей воды не заданы.',
15:'Две NET, третий RFQ. Два чиллера сохранены; тепловая нагрузка, климат помещений и состав управления открыты.',
16:'Один ценовой ориентир насоса. Нужны циркуляция, морская вода, клапаны, трубы, изоляция и полная автоматика.',
17:'Смесительный узел и фильтр — фрагменты. Количества, аккумуляторы давления и полная обвязка открыты.',
18:'Одна меняющаяся цена duplex-фильтра. Нужны вся топливная схема, обратки, сепарация и пожарное отсечение.',
19:'Цены фильтров и отдельные части генераторного пакета. Водозаборы, охлаждение, выхлопы и длины открыты.',
20:'Цены насоса, ручного насоса и сигнализации не закрывают систему. Нужны расчёт притока, напора и отсеков.',
21:'Поисковая цена исключена. Нужны защищаемый объём, агент тушения, отключения и полное согласованное оснащение.',
22:'Цена не допущена; напор/расход и тепловыделение не заданы. Вентиляторы, каналы и заслонки открыты.',
23:'Есть насос замены масла; свет, сервисные соединения и инструментальная обвязка открыты.',
24:'Есть только датчики, не танки. Объёмы, встроенное изготовление, люки и испытания открыты.',
25:'Один насос стоков; танки, вентиляция, управление и полный маршрут открыты.',
26:'Опция. Один NET комплекта котла, другой налоговый конфликт исключён; полный отопительный контур не оценён.',
27:'Две АКБ сохранены; выбранные генераторы, пусковой ток и напряжение открыты, цены нет.',
28:'Два контактора как разные компоненты; не полная резервная система. Нужны изоляция отказавшей АКБ и совместимые пути пуска.',
29:'Три архитектуры, не три готовые станции. PTO/резерв, тяга подрулек, лебёдки, одновременность и полный состав открыты.'}
coverage=[]
for m in read(R/'deliverables/M01/v001/01_support/equipment-list.csv'):
    candidates=[r for r in rows if m['item_id'] in r['me_id'].split('|')]
    coverage.append({'me_id':m['item_id'],'equipment':m['equipment'],'quantity':m['quantity'],
      'quantity_status':m['quantity_status'],'records':len(candidates),
      'net_records':sum(bool(r['net_unit_price']) for r in candidates),
      'candidate_ids':' | '.join(r['candidate_id'] for r in candidates),
      'complete_three_offers':'NO — OUTPUT/SCOPE OPEN','next_work':gaps[int(m['item_id'][-2:])]})
assert len(coverage)==29
write(S/'coverage.csv',coverage)

# Price order is a research aid, never a quality rating or an approved basket.
groups=[
 ('Генераторы, 2 шт','PROP-G03 PROP-G01 PROP-G02','Три цены за разные исполнения; нагрузка и резерв при отказе ещё не проверены.'),
 ('Инверторы, 2 шт','EL-INV-A EL-INV-B EL-INV-C','Мощность поиска, не проект. Различия зарядки/мониторинга/схемы доводятся до нужного результата.'),
 ('Сервисные АКБ, 6 шт условно','EL-BAT-A EL-BAT-B EL-BAT-C','Только ветвь 115 А·ч / 12 В; нужны подтверждение владельца, доступная энергия и ресурс.'),
 ('Бойлер, цена за шт','WC-HW-01 WC-HW-02 WC-HW-03','40 л — пример, количество открыто. Сравнить полезный запас горячей воды и обслуживание.'),
 ('Чиллеры, 2 шт','WC-CH-01 WC-CH-02 WC-CH-03','16 000 BTU/h — пример, не выбранная мощность. Панели/насосы/воздушные блоки уточнять.'),
 ('Опреснитель, 1 шт','WC-RO-01 WC-RO-02 WC-RO-03','Базовая цена и цена ОТ не равны полному результату 150 л/ч с автоматикой.'),
 ('Пресная вода 24 В, 1 установка','WC-DC-01 WC-DC-02 WC-DC-03','Сопоставимость рабочей точки и состав открыты.'),
 ('Пресная вода AC, 1 установка','WC-AC-01 WC-AC-02 WC-AC-03','Одна открытая цена. Неизвестные не ранжировать как дешёвые.'),
 ('Развязка берега, цена за шт','EL-ISO-A EL-ISO-B EL-ISO-C','Выбрать по нужной береговой мощности, потерям и месту установки; количество открыто.'),
 ('Двигатель с редуктором, 1 комплект','PROP-E01 PROP-E02 PROP-E03','ТОЛЬКО исторический март 2023; не текущий диапазон цен 2026.'),
 ('Гидравлика, архитектуры','A01 A07 A08','Vetus — цена бака, не станции; Sleipner/Harken RFQ. Тяга и работа лебёдок открыты.')]
comparison=[]
for title,ids,limit in groups:
    rs=[byid[x] for x in ids.split()]
    sortable=[r for r in rs if r['eur_net_unit']]
    ordered=sorted(sortable,key=lambda r:D(r['eur_net_unit']))
    for r in rs:
        rank='Без ранга: нет NET'
        if r in ordered:
            n=ordered.index(r)
            rank=(['Нижняя','Средняя','Верхняя'][n] if len(ordered)==3 else f'{n+1} из {len(ordered)} известных')
        if title.startswith('Гидравлика'): rank='Часть архитектуры, не ранг'
        comparison.append({'group':title,'candidate_id':r['candidate_id'],'model':r['model'],'price_order_only':rank,
          'eur_net_unit':r['eur_net_unit'],'quantity':r['qty'],'eur_fixed_qty_line':r['eur_fixed_qty_line'],
          'price_class':r['price_class'],'limitation':limit,'price_url':r['price_url']})
write(S/'comparison.csv',comparison)

dimensions=[]
for r in rows:
    ds=[r[x] for x in ['length_mm','width_mm','height_mm']]
    if all(ds):
        l,w,h=map(D,ds)
        dimensions.append({'candidate_id':r['candidate_id'],'model':r['model'],'length_mm':ds[0],'width_mm':ds[1],
         'height_mm':ds[2],'mass_kg':r['mass_kg'],'box_m3':str((l*w*h/D(10**9)).quantize(D('.0001'))),
         'installation_limit':'Габарит изделия/указанного модуля, не сервисный объём. Ориентация и состав: см. примечание.',
         'scope_notes':r['scope_notes'],'spec_url':r['spec_url']})
write(S/'dimensions.csv',dimensions)

wb=Workbook(); wb.remove(wb.active)
def sheet(name, headers, data, widths=None):
    ws=wb.create_sheet(name); ws.append(headers)
    for row in data:
        ws.append([float(v) if isinstance(v,D) else v for v in row])
    ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
    ws.sheet_view.zoomScale=80
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor='17384D'); c.font=Font(color='FFFFFF',bold=True)
        c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=34
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if c.row%2==0: c.fill=PatternFill('solid',fgColor='F0F5F8')
            if isinstance(c.value,str) and c.value.startswith('https://'):
                c.hyperlink=c.value; c.font=Font(color='006CA8',underline='single')
            if isinstance(c.value,(float,int)): c.number_format='#,##0.00'
        ws.row_dimensions[row[0].row].height=78
    for i in range(1,len(headers)+1): ws.column_dimensions[get_column_letter(i)].width=(widths or {}).get(i,23)
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:1'
    return ws
def num(x): return D(x) if x not in ['',None] else None
summary=[
 ['Статус','M02 / 11.09.2026 — рыночный анализ с открытыми входами; не КП, не готовая КАЦ и не бюджет МО.'],
 ['Объём',f'{len(rows)} записей: {sum(bool(r["net_unit_price"]) for r in rows)} с числом NET, включая исторические/ОТ/нестабильные; 29 групп исходного перечня.'],
 ['Принцип владельца','Сравниваем полезный результат. Разные марки, номиналы и состав обвязки допустимы при нужных функциях и фиксированных количествах основных устройств.'],
 ['Зафиксировано','Только 24 В DC / 220 В AC; все подруливающие и лебёдки гидравлические; 6 сервисных + 2 генераторных + 1 пусковая АКБ.'],
 ['Не задано','Рабочие нагрузки, пусковой ток, климат помещений, рабочие точки воды/гидравлики, 115 ампер/А·ч и напряжение сервисных АКБ.'],
 ['Как читать цены','Нижняя/средняя/верхняя — порядок известных цен, не оценка качества. Полные три функционально сопоставимые поставки пока не сформированы.'],
 ['Что включено','Только указанный товар/комплект. Доставка в Польшу, импорт, монтаж, кабели/трубы, опоры и испытания не включены, если прямо не оговорены источником.'],
 ['Валюта и арифметика','EUR для сравнения; GBP / 0.85915 по ECB 10.09.2026. NET единицы округлён до цента HALF_UP, затем EUR до цента, затем × количество. Суммы альтернатив не складывать.'],
 ['Цена и дата','11.09.2026 — дата просмотра. Дата действия большинства цен неизвестна. Volvo — март 2023. Цена ОТ и колеблющаяся цена выделены.'],
 ['Компактность','Размер корпуса изделия не равен месту установки: нужны доступ к сервису, охлаждение, кабельные радиусы и путь извлечения. Помещаемость не доказана.'],
 ['Следующее действие','Задать общие выходные критерии и наложить крупные агрегаты с сервисными зонами на измеренное МО; затем довести три варианта до полной поставки и КП.']]
sheet('00_Как читать',['Тема','Вывод'],summary,{1:27,2:130})
sheet('01_Сравнение',['Группа','ID','Модель','Порядок цены','EUR NET/ед.','Кол-во','EUR NET строки','Класс цены','Ограничение','Источник цены'],
 [[r['group'],r['candidate_id'],r['model'],r['price_order_only'],num(r['eur_net_unit']),num(r['quantity']),num(r['eur_fixed_qty_line']),r['price_class'],r['limitation'],r['price_url']] for r in comparison],{1:31,3:42,8:35,9:70,10:45})
fields=['candidate_id','me_id','manufacturer','model','qty','unit','currency','listed_price','tax_basis','tax_rate','net_unit_price','eur_net_unit','eur_fixed_qty_line','price_class','technical_status','scope_notes','availability','price_url','spec_url','access_date','price_date','source_register']
headers=['ID','Пакет M01','Производитель','Модель','Кол-во','Ед.','Валюта','Цена источника','Налоговая основа','Ставка %','NET/ед. в валюте','EUR NET/ед.','EUR NET строки','Класс цены','Технический статус','Состав и ограничения','Наличие','Источник цены','Спецификация','Просмотр','Дата цены','Реестр специалиста']
numeric={'qty','listed_price','tax_rate','net_unit_price','eur_net_unit','eur_fixed_qty_line'}
sheet('02_Все кандидаты',headers,[[num(r[k]) if k in numeric else r[k] for k in fields] for r in rows],{4:45,9:38,14:38,15:48,16:100,17:40,18:45,19:45})
sheet('03_Покрытие',['M01','Оборудование','Кол-во M01','Статус кол-ва','Записей','С NET','Кандидаты','Три полных варианта','Что ещё требуется'],
 [[r[k] for k in coverage[0]] for r in coverage],{2:40,7:60,8:35,9:100})
sheet('04_Габариты',['ID','Модель','Размер 1 мм','Размер 2 мм','Размер 3 мм','Масса кг','Коробка м³','Предел применения','Состав / сервис','Первичный источник'],
 [[r['candidate_id'],r['model'],num(r['length_mm']),num(r['width_mm']),num(r['height_mm']),num(r['mass_kg']),num(r['box_m3']),r['installation_limit'],r['scope_notes'],r['spec_url']] for r in dimensions],{2:45,8:60,9:100,10:45})
for c in wb['04_Габариты']['G'][1:]: c.number_format='0.0000'
outcomes=read(S/'service-outcomes.csv')
sheet('05_Полезный результат',['Пакеты','Система','Требуемый результат','Фиксированное количество','Открытые показатели','Критерии сравнения'],[[r[k] for k in outcomes[0]] for r in outcomes],{1:30,2:30,3:80,4:43,5:70,6:85})
sheet('06_Правила и курс',['Параметр','Значение'],[
 ['Источник GBP/EUR',json.loads((S/'fx.json').read_text())['source']],['Курс','1 EUR = 0.85915 GBP; наблюдение 10.09.2026'],
 ['Неизвестное','Пустая ячейка означает неизвестно; это не нулевая стоимость или масса.'],
 ['КП','Ни одна публичная карточка не объявлена полученным коммерческим предложением.'],
 ['НДС','NET выделен по прямой маркировке либо подтверждённой ставке исходного продавца; налоговая схема Польши не назначена.'],
 ['Двойной учёт','A06 HT1013 включён в A01; не добавлять. Редуктор в PROP-E01/02/03 не считать второй раз.'],
 ['Резервный пуск','Сначала изолировать отказавшую АКБ. Проверить напряжение, CCA, проводку и коммутацию каждого резервного пути. Работа от любой одной АКБ ещё не доказана.'],
 ['Округление','HALF_UP: NET единицы → 2 знака; GBP в EUR → 2 знака; отображённая EUR цена × количество → 2 знака. Это расчётный ориентир, не счёт продавца.'],
 ['УТВ','Пустой шаблон U01 не заполнен и не изменён.']],{1:28,2:135})
target=S.parent/'00_market-analysis.xlsx'; wb.save(target)
check=load_workbook(target,read_only=True,data_only=True)
assert len(check.sheetnames)==7 and check['02_Все кандидаты'].max_row==66
assert check['03_Покрытие'].max_row==30
check.close()
print(json.dumps({'records':len(rows),'with_net':sum(bool(r['net_unit_price']) for r in rows),'coverage':len(coverage),'dimensions':len(dimensions),'workbook':str(target)},ensure_ascii=False))
