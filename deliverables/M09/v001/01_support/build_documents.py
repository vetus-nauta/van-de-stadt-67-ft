"""M09: one source register and BOM drive both the revised KAC and estimate.

Old sources retain their observation dates. New offers are not supplier quotations.
Selection/quantities are explicit planning assumptions in integration.json.
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv, json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

S=Path(__file__).resolve().parent
R=S.parents[3]
D=Decimal
def cash(x): return D(str(x)).quantize(D('.01'),rounding=ROUND_HALF_UP)
def number(x): return float(x) if x not in ('',None) else None
def read(path): return list(csv.DictReader((R/path).open(),delimiter=';'))
def save_csv(name,rows):
    with (S/name).open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]),delimiter=';',lineterminator='\n'); w.writeheader();w.writerows(rows)
def text(x): return json.dumps(x,ensure_ascii=False) if isinstance(x,(dict,list)) else str(x or '')
cfg=json.loads((S/'integration.json').read_text())
oldp=read('deliverables/M03/v001/01_support/price-register.csv')
oldc=read('deliverables/M03/v001/01_support/components.csv')
packages=read('deliverables/M03/v001/01_support/packages.csv')
sources={}
for p in oldp:
    sources[p['candidate_id']]=dict(id=p['candidate_id'],me=p['me_id'],name=p['manufacturer']+' '+p['model'],sku=p['model'],seller='См. исходную карточку M03',url=p['price_url'],viewed_on=p['access_date'] or '2026-09-11',currency=p['currency'],source_price=p['listed_price'],source_tax_percent=p['tax_rate'],source_tax_basis=p['tax_basis'],net_unit_eur=p['eur_net_unit'],unit=p['unit'],scope_included=p['scope_notes'],scope_excluded='Исходные ограничения M03 сохранены',compatibility=p['technical_status'],price_evidence=p['evidence_type'],revision='M03/v001; цена не перепроверялась в M09',fx_date='2026-09-10' if p['currency']=='GBP' else '',fx_quote_per_eur='0.85915' if p['currency']=='GBP' else '')
new=[]
for path in sorted(S.glob('*/offers.json')):
    for o in json.loads(path.read_text()):
        o=dict(o);o['me']=re.sub(r'ME-?(\d{2})',r'ME-\1',o['me']);o['revision']=str(path.relative_to(R))
        assert o['id'] not in sources,o['id']
        sources[o['id']]=o;new.append(o)

components=[]
for c in oldc:
    if c['variant_id'].startswith('PK12-'): continue
    original=c['candidate_id']; cid=cfg['replace'].get(original,original)
    p=sources[cid]
    q=c['quantity']; qb='Количество исходной M03; основные агрегаты без изменения'
    if q=='':
        setting=cfg['quantities'].get(original)
        if setting:q,qb=setting['qty'],setting['basis']
    components.append(dict(id=c['line_id'],variant=c['variant_id'],package=c['variant_id'][:4],me=c['me_id'],source=cid,qty=q,unit=c['unit'],qty_basis=qb,note=('Подстановка M09 вместо '+original+'. Прежнее описание относится к старому SKU и сохранено в M03.' if cid != original else c['scope_note']),origin=c['line_id'],optional=False))
for extra in cfg['additions']:
    for letter in 'ABC':
        cid=extra['source'];p=sources[cid]
        components.append(dict(id=extra['id']+'-'+letter,variant=extra['package']+'-'+letter,package=extra['package'],me=extra['me'],source=cid,qty=extra.get('qty'),unit=extra.get('unit',p.get('unit','шт')),qty_basis=extra['basis'],note=extra['note'],origin=extra.get('origin','NEW'),optional=extra.get('optional',False)))

def make_row(c):
    p=sources[c['source']]; net=p.get('net_unit_eur'); gross_unit=58842 if p['id']=='YAN-E01' else None
    if p['id'] in cfg.get('gross_only',{}):gross_unit=cfg['gross_only'][p['id']]
    mode='SOURCE_GROSS' if gross_unit is not None else 'PL23_ESTIMATE'
    qty=c['qty']; ext=cash(D(str(qty))*D(str(net))) if qty not in ('',None) and net not in ('',None) else None
    vat=cash(ext*D('.23')) if ext is not None else None
    gross=cash(D(str(qty))*D(str(gross_unit))) if gross_unit is not None and qty not in ('',None) else (ext+vat if ext is not None else None)
    return dict(id=c['id'],package=c['package'],me=c['me'],type='Опция' if c['optional'] else 'Оборудование / компонент',qty=qty,unit=c['unit'],name=p['name'],source=p['id'],net_unit=net,tax_mode=mode,rate=None if gross_unit is not None else 23,source_gross_eur=gross_unit,scope_status='ОПЦИЯ' if c['optional'] else 'СОСТАВ И СОВМЕСТИМОСТЬ ОТКРЫТЫ',price_class='ПУБЛИЧНАЯ ЦЕНА / ПРЕДВАРИТЕЛЬНАЯ ПОДСТАНОВКА' if p in new else 'ПЕРЕНОС M03, ДАТА СОХРАНЕНА',note=c['note']+' | '+text(p.get('scope_included'))+' | Исключено: '+text(p.get('scope_excluded'))+' | '+text(p.get('compatibility')),url=p['url'],original_price=p.get('source_price'),original_currency=p.get('currency'),net_eur=str(ext) if ext is not None else '',vat_eur=str(vat) if vat is not None else '',gross_eur=str(gross) if gross is not None else '',qty_basis=c['qty_basis'],observed=p.get('viewed_on',''),origin=c['origin'],included=not c['optional'])

allrows=[make_row(c) for c in components]
rows=[r for c,r in zip(components,allrows) if c['variant'].endswith('-B')]
oldestimate=read('deliverables/M07/v001/01_support/estimate-lines.csv')
for old in oldestimate:
    if not (old['id'].startswith('GAP-') or old['id'] in cfg['retain_open']):continue
    r=dict(old);r.update(qty_basis='Требуются проект/объёмы/комплектность или условия поставки',observed='',origin=old['id'],included=True)
    if old['id'] in cfg.get('open_notes',{}):r['note']=cfg['open_notes'][old['id']]
    rows.append(r)
for extra in cfg.get('open_residuals',[]):
    r={k:'' for k in rows[0]};r.update(extra);r.update(type='Неоценённый остаток системы',scope_status='ОТКРЫТО',price_class='ТРЕБУЕТСЯ ПРОЕКТ / КП',included=True,tax_mode='PL23_ESTIMATE',rate=23);rows.append(r)
save_csv('estimate-lines.csv',rows)
save_csv('components.csv',components)
save_csv('calculated-components.csv',[dict(variant=c['variant'],**r) for c,r in zip(components,allrows)])
keys=['id','me','name','sku','seller','url','viewed_on','currency','source_price','source_tax_basis','source_tax_percent','net_unit_eur','unit','scope_included','scope_excluded','compatibility','price_evidence','revision','fx_date','fx_quote_per_eur']
source_rows=[{k:text(p.get(k)) for k in keys} for p in sources.values()]
save_csv('price-register.csv',source_rows)
included=[r for r in rows if r['included']]
summoney=lambda key,items:sum((D(r[key]) for r in items if r[key] not in ('',None)),D(0))
stats={'main_rows':len(rows),'priced_gross_rows':sum(bool(r['gross_eur']) for r in included),'unpriced_gross_rows':sum(not bool(r['gross_eur']) for r in included),'priced_net_rows':sum(bool(r['net_eur']) for r in included),'known_net_eur':str(summoney('net_eur',included)),'estimated_vat_identified_part_eur':str(summoney('vat_eur',included)),'source_gross_tax_not_split_eur':str(sum((D(r['gross_eur']) for r in included if r['tax_mode']=='SOURCE_GROSS' and r['gross_eur']),D(0))),'known_gross_mixed_basis_eur':str(summoney('gross_eur',included)),'optional_gross_eur':str(summoney('gross_eur',[r for r in rows if not r['included']])),'new_price_records':len(new),'baseline_gross_eur':'131819.50','status':'PLANNING_COMPONENT_COST_NOT_COMPLETE_INSTALLED_BUDGET'}
stats['change_gross_eur']=str(D(stats['known_gross_mixed_basis_eur'])-D(stats['baseline_gross_eur']))
(S/'totals.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2)+'\n')

def sheet(wb,name,headers,data,widths=None):
    ws=wb.create_sheet(name);ws.append(headers)
    for r in data:ws.append(r)
    ws.freeze_panes='E2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.zoomScale=75
    for row in ws:
        ws.row_dimensions[row[0].row].height=70 if row[0].row>1 else 45
        for c in row:
            c.font=Font(name='Calibri',size=11);c.alignment=Alignment(wrap_text=True,vertical='top')
            if c.row==1:c.fill=PatternFill('solid',fgColor='234351');c.font=Font(name='Calibri',color='FFFFFF',bold=True)
            elif c.row%2==0:c.fill=PatternFill('solid',fgColor='F1F5F7')
            if c.data_type=='f':c.number_format='#,##0.00';c.fill=PatternFill('solid',fgColor='E1ECEF')
            if isinstance(c.value,str) and c.value.startswith('http'):c.hyperlink=c.value;c.font=Font(color='126882',underline='single')
    for i in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(i)].width=(widths or {}).get(i,23)
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
    return ws
def book():
    wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True);return wb

notes=[['Статус','M09/v001 • 15.09.2026. Дополнение ценами SVB и других морских магазинов. Это рабочая оценка компонентов, не полная стоимость установленного МО и не утверждённая УТВ.'],['Количество','1 двигатель с 1 редуктором, 2 генератора, 2 основных инвертора, 6+2+1 АКБ сохранены. Ранее пустые количества отдельных компонентов заданы только в явно обозначенном предварительном сценарии.'],['Исходные цены','Новые источники датированы 15.09.2026; перенесённые цены M03 сохраняют прежнюю дату и ограничения. Ни одна карточка магазина не названа полученным персональным КП.'],['НДС','NET извлечён по налоговому отображению конкретной карточки. В основной смете к NET добавлены предварительные 23% по D030. SOURCE_GROSS не облагается повторно; его ставка неизвестна.'],['GBP','Новые цены GBP переведены по ЕЦБ 14.09.2026: 1 EUR = 0.85598 GBP. Старые цены M03 с прежним FX 10.09 сохранены и маркированы. Импорт UK/доставка/таможня требуют отдельной оценки.'],['Совместимость','12V стартерные кандидаты — только ценовые аналоги; требование владельца 24 V не изменено. Полная схема пуска от любого АКБ не разработана. PTO/Yanmar, объём пожаротушения, диаметры и расходы не подтверждены ценой.'],['Состав','Дополнительные компоненты не объявлены полными системами. За составом остались отдельные неоценённые строки; готовые баки Gianneschi 20 L и клапан HT1013 внутри гидробака повторно не начислены.'],['Опции','Резервное водяное отопление остаётся опцией и показано отдельно от основного итога.'],['A/B/C','Сохранены альтернативы M03, без ранжирования качества по цене. Общие новые компоненты применены с одинаковыми количествами; общий ценовой аналог не выдаётся за три независимых предложения.'],['Изменения','Исходные M03/M07 и комплект 1.0 сохранены. Рост суммы включает раскрытие состава и предварительных количеств, а не только изменение рыночных цен.']]

wb=book()
overview=sheet(wb,'00_Итог',['Показатель','Значение','Основание'],[
 ['Оценённая часть NET EUR',number(stats['known_net_eur']),'Без строк SOURCE_GROSS'],['Предварительный НДС EUR',number(stats['estimated_vat_identified_part_eur']),'23% построчно'],['Источник с невыделенным налогом EUR',number(stats['source_gross_tax_not_split_eur']),'Без повторного начисления'],['Оценённая часть с налогом EUR',number(stats['known_gross_mixed_basis_eur']),'Смешанная налоговая основа; неполный бюджет'],['Опциональное отопление с налогом EUR',number(stats['optional_gross_eur']),'Отдельно, не включено'],['Основных строк с суммой',stats['priced_gross_rows'],'Компонентная оценка'],['Основных строк без суммы',stats['unpriced_gross_rows'],'Неизвестное не равно нулю'],['Изменение против M07 EUR',number(stats['change_gross_eur']),'Добавлены цены, количества и состав']]+[[a,b,''] for a,b in notes],{1:42,2:100,3:75})
data=[]
for i,r in enumerate(rows,2):
    net=f'=IF(AND(ISNUMBER(E{i}),ISNUMBER(I{i})),ROUND(E{i}*I{i},2),"")'
    vat=f'=IF(AND(K{i}="PL23_ESTIMATE",ISNUMBER(J{i}),ISNUMBER(L{i})),ROUND(J{i}*L{i}/100,2),"")'
    gross=f'=IF(K{i}="SOURCE_GROSS",IF(AND(ISNUMBER(E{i}),ISNUMBER(N{i})),ROUND(E{i}*N{i},2),""),IF(AND(ISNUMBER(J{i}),ISNUMBER(M{i})),ROUND(J{i}+M{i},2),""))'
    data.append([r['id'],r['package'],r['me'],r['type'],number(r['qty']),r['unit'],r['name'],r['source'],number(r['net_unit']),net,r['tax_mode'],number(r['rate']),vat,number(r['source_gross_eur']),gross,r['scope_status'],r['price_class'],r['note'],r['url'],number(r['original_price']),r['original_currency'],r['qty_basis'],r['observed'],'ДА' if r['included'] else 'ОПЦИЯ',r['origin']])
ws=sheet(wb,'01_Смета',['ID','Пакет','Группы','Вид','Кол-во','Ед.','Наименование','Источник','Цена NET EUR','Сумма NET EUR','Налоговая основа','НДС % план','НДС EUR','Исходный gross EUR','С налогом EUR','Комплектность','Статус цены','Состав / пределы','Прямая ссылка','Цена источника','Валюта','Основание количества','Дата просмотра','В основном итоге','Связь с M07'],data,{7:55,18:95,19:48,22:85})
last=ws.max_row;z=last+1;ws.cell(z,7,'Оценённая часть основного сценария — НЕ полный бюджет')
for col in ['J','M','O']:ws[f'{col}{z}']=f'=SUMIF(X2:X{last},"ДА",{col}2:{col}{last})';ws[f'{col}{z}'].number_format='#,##0.00'
overview['B2']=f"='01_Смета'!J{z}";overview['B3']=f"='01_Смета'!M{z}";overview['B5']=f"='01_Смета'!O{z}"
sheet(wb,'02_Источники',keys,[[p[k] for k in keys] for p in source_rows],{3:55,6:60,14:90,15:90,16:90})
sheet(wb,'03_Изменения',['Позиция M07','Действие / предположение'],[[k,v] for k,v in cfg['change_notes'].items()],{1:32,2:145})
sheet(wb,'04_Открыто',['ID','Группа','Открытая часть','Дальнейший вход'],[[r['id'],r['me'],r['name'],r['note']] for r in rows if not r['gross_eur']],{3:65,4:145})
sheet(wb,'05_Основания',['Тема','Условие'],notes,{1:35,2:150})
wb.save(S.parent/'00_engine-room-estimate.xlsx')

kac=book()
sheet(kac,'00_Основания',['Тема','Условие'],notes,{1:35,2:150})
cdata=[]
for i,(c,r) in enumerate(zip(components,allrows),2):
    cdata.append([c['variant'],c['id'],c['me'],r['name'],number(r['qty']),r['unit'],r['source'],number(r['net_unit']),f'=IF(AND(ISNUMBER(E{i}),ISNUMBER(H{i})),ROUND(E{i}*H{i},2),"")',r['tax_mode'],number(r['rate']),f'=IF(AND(J{i}="PL23_ESTIMATE",ISNUMBER(I{i})),ROUND(I{i}*K{i}/100,2),"")',number(r['source_gross_eur']),f'=IF(J{i}="SOURCE_GROSS",IF(AND(ISNUMBER(E{i}),ISNUMBER(M{i})),ROUND(E{i}*M{i},2),""),IF(ISNUMBER(I{i}),ROUND(I{i}+L{i},2),""))',r['qty_basis'],r['note'],r['observed'],r['url'],'ОПЦИЯ' if c['optional'] else 'ДА'])
cs=sheet(kac,'02_Состав',['Вариант','ID','ME','Исполнение','Кол-во','Ед.','Источник','NET EUR/ед.','NET EUR','Налог','План %','НДС EUR','Исх.gross/ед.','С налогом EUR','Основание количества','Открытый состав / совместимость','Дата','URL','В основном итоге'],cdata,{4:65,15:85,16:105,18:55})
end=cs.max_row;kdata=[]
for p in packages:
    for l in 'ABC':
        v=p['package_id']+'-'+l;i=len(kdata)+2
        active=[r for c,r in zip(components,allrows) if c['variant']==v and not c['optional']]
        kdata.append([v,p['title'],p['fixed_quantities'],f'=SUMIFS(\'02_Состав\'!I2:I{end},\'02_Состав\'!A2:A{end},A{i},\'02_Состав\'!S2:S{end},"ДА")',f'=SUMIFS(\'02_Состав\'!L2:L{end},\'02_Состав\'!A2:A{end},A{i},\'02_Состав\'!S2:S{end},"ДА")',f'=SUMIFS(\'02_Состав\'!N2:N{end},\'02_Состав\'!A2:A{end},A{i},\'02_Состав\'!S2:S{end},"ДА")',sum(not bool(r['gross_eur']) for r in active),'НЕПОЛНЫЙ ПАКЕТ',p['scope_to_complete'],p['open_inputs']])
ks=sheet(kac,'01_КАЦ',['Вариант','Система','Фиксированные количества','Оценено NET EUR','Плановый налог EUR','Оценено с налогом EUR','Строк без суммы','Статус','Полный требуемый состав','Открытые входы'],kdata,{2:42,3:60,9:100,10:100})
for l in 'ABC':
    z=ks.max_row+1;ks.cell(z,1,'Итого '+l);ks.cell(z,2,'Частичная компонентная оценка, не полная стоимость установки')
    for col in 'DEF':ks[f'{col}{z}']=f'=SUMIF(A2:A37,"*-{l}",{col}2:{col}37)';ks[f'{col}{z}'].number_format='#,##0.00'
sheet(kac,'03_Все цены',keys,[[p[k] for k in keys] for p in source_rows],{3:55,6:60,14:90,15:90,16:90})
newids={p['id'] for p in new}
sheet(kac,'04_Новые предложения',keys,[[p[k] for k in keys] for p in source_rows if p['id'] in newids],{3:55,6:60,14:90,15:90,16:90})
sheet(kac,'05_Неоценённый остаток',['ID','Группа','Затрата','Почему не закрыта магазинной ценой'],[[r['id'],r['me'],r['name'],r['note']] for r in rows if not r['gross_eur']],{3:60,4:145})
scope=read('deliverables/M06/v001/01_support/scope-register.csv')
sheet(kac,'06_Комплектность_M06',list(scope[0]),[[r[k] for k in scope[0]] for r in scope],{3:55,10:70,13:80,14:80})
kac._sheets.sort(key=lambda w:w.title)
kac.save(S.parent/'02_working-KAC.xlsx')
print(json.dumps(stats,ensure_ascii=False,indent=2))

prices=book()
sheet(prices,'00_Условия',['Тема','Условие'],notes,{1:35,2:150})
sheet(prices,'01_Новые цены',keys,[[p[k] for k in keys] for p in source_rows if p['id'] in newids],{3:55,6:60,14:90,15:90,16:90})
sheet(prices,'02_Все источники',keys,[[p[k] for k in keys] for p in source_rows],{3:55,6:60,14:90,15:90,16:90})
prices.save(S.parent/'03_public-price-offers.xlsx')
